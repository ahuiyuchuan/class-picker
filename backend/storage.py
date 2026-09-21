"""SQLite 本地业务服务。

@author ahui
同一数据目录的操作通过可重入进程互斥串行，连接按请求创建；结果和进度同事务保存。
数据库不存入 EXE 解包目录。测试通过构造函数传入隔离目录，不触碰真实数据。
"""
from contextlib import contextmanager, closing
import csv
import io
import json
import secrets
import sqlite3
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from data_lock import DataLock

DEFAULTS = dict(count=1, mode='unique', effect='rolling', duration=1,
                show_number=True, font_size='auto', start_fullscreen=False,
                save_history=True, history_days=90, auto_backup=True, backup_count=7)
RULE_KEYS = {'count', 'mode', 'effect', 'duration', 'show_number'}


def valid_settings(values, partial=False):
    """拒绝未知字段及错误类型；班级只允许覆盖明确支持的规则。"""
    if not isinstance(values, dict):
        raise ValueError('设置格式无效')
    allowed = RULE_KEYS if partial else DEFAULTS.keys()
    if set(values) - set(allowed):
        raise ValueError('包含不支持的设置')
    result = dict(values) if partial else {**DEFAULTS, **values}
    choices = dict(mode=['unique', 'random'], effect=['direct', 'reveal', 'rolling'],
                   duration=[.5, 1, 2], font_size=['auto', 'normal', 'large', 'huge'],
                   history_days=[30, 90, 0], backup_count=[3, 7, 15])
    for key, value in result.items():
        if key in ('show_number', 'start_fullscreen', 'save_history', 'auto_backup'):
            if type(value) is not bool:
                raise ValueError('开关必须为布尔值')
        elif key == 'count':
            if type(value) is not int or not 1 <= value <= 100:
                raise ValueError('抽取人数须为 1—100 的整数')
        elif key in choices and (isinstance(value, bool) or value not in choices[key]):
            raise ValueError('设置值超出允许范围')
    return result


class Store:
    """负责班级、名单、规则和备份；调用方只传业务参数，不接受 SQL。"""

    def __init__(self, folder):
        self.folder = Path(folder)
        self.folder.mkdir(parents=True, exist_ok=True)
        self.path = self.folder / 'app.db'
        self.backups = self.folder.parent / 'backups'
        self.backups.mkdir(exist_ok=True)
        self.lock = DataLock(self.path)
        with self.lock, self.connect() as db:
            version = db.execute('PRAGMA user_version').fetchone()[0]
            if version > 2:
                raise ValueError('数据库版本较新，请使用新版程序')
            if version == 1:
                # 升级前保留旧版可读取的副本，回退程序时可恢复该副本。
                self.backup(kind='pre-v2')
            db.executescript('''
                CREATE TABLE IF NOT EXISTS classroom (
                    id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE,
                    position INTEGER NOT NULL DEFAULT 0, rules TEXT NOT NULL DEFAULT '{}',
                    latest TEXT NOT NULL DEFAULT '[]', request TEXT);
                CREATE TABLE IF NOT EXISTS student (
                    id INTEGER PRIMARY KEY, class_id INTEGER NOT NULL REFERENCES classroom(id) ON DELETE CASCADE,
                    name TEXT NOT NULL, no TEXT, enabled INTEGER NOT NULL DEFAULT 1,
                    drawn INTEGER NOT NULL DEFAULT 0, excluded INTEGER NOT NULL DEFAULT 0,
                    UNIQUE(class_id,no));
                CREATE INDEX IF NOT EXISTS student_class ON student(class_id);
                CREATE TABLE IF NOT EXISTS draw_record (
                    id INTEGER PRIMARY KEY, class_id INTEGER NOT NULL REFERENCES classroom(id) ON DELETE CASCADE,
                    request TEXT NOT NULL UNIQUE, time TEXT NOT NULL, results TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS record_class_time ON draw_record(class_id,time);
                CREATE TABLE IF NOT EXISTS app_setting (key TEXT PRIMARY KEY, value TEXT NOT NULL);
            ''')
            self.migrate_student_order(db)

    @staticmethod
    def migrate_student_order(db):
        """将旧名单按原 ID 顺序迁移为可排序名单；启动和恢复旧备份均调用。

        position 仅控制展示/导出顺序，不改变学生 ID、抽取状态或历史。
        新插入记录由触发器追加到本班末尾；迁移可重复执行，不重置已有排序。
        """
        columns = {row[1] for row in db.execute('PRAGMA table_info(student)')}
        if 'position' not in columns:
            db.execute('ALTER TABLE student ADD COLUMN position INTEGER NOT NULL DEFAULT 0')
            db.execute('UPDATE student SET position=id')
        # 覆盖追加 MAX(position) 和名单排序，避免批量导入每一行扫描整个班级。
        db.execute('CREATE INDEX IF NOT EXISTS student_class_position ON student(class_id,position)')
        # 历史按 ID 倒序展示；旧的 class_id,time 索引不能直接满足该排序。
        db.execute('CREATE INDEX IF NOT EXISTS record_class_id ON draw_record(class_id,id)')
        # 升级/恢复时重建程序自有触发器，不沿用备份中可能被修改的同名 SQL。
        db.execute('DROP TRIGGER IF EXISTS student_append_position')
        db.execute('''CREATE TRIGGER student_append_position
            AFTER INSERT ON student BEGIN
                UPDATE student SET position=(SELECT COALESCE(MAX(position),0)+1
                    FROM student WHERE class_id=NEW.class_id AND id<>NEW.id)
                WHERE id=NEW.id;
            END''')
        db.execute('PRAGMA user_version=2')

    @contextmanager
    def connect(self):
        """每次独立连接，外键生效，最长等待写锁 5 秒。"""
        db = sqlite3.connect(self.path, timeout=5)
        db.row_factory = sqlite3.Row
        db.execute('PRAGMA foreign_keys=ON')
        try:
            with db:
                yield db
        finally:
            db.close()

    def config(self, db):
        row = db.execute("SELECT value FROM app_setting WHERE key='config'").fetchone()
        return valid_settings(json.loads(row[0]) if row else {})

    def snapshot(self):
        """在同一读取事务中用五次查询组装完整快照，避免班级数带来的 N+1。

        名单批量读取，历史通过班级索引各取最近 50 条；不加载全部历史。
        BEGIN 保证独立进程写入时，班级、名单、规则和统计来自同一数据库版本。
        """
        with self.lock, self.connect() as db:
            db.execute('BEGIN')
            classes = []
            for row in db.execute('SELECT * FROM classroom ORDER BY position,id'):
                item = dict(row)
                item['rules'] = json.loads(item['rules'])
                item['latest'] = json.loads(item['latest'])
                item.pop('request')
                item['students'] = []
                item['history'] = []
                classes.append(item)
            by_id = {item['id']: item for item in classes}
            for row in db.execute('SELECT * FROM student ORDER BY class_id,position,id'):
                by_id[row['class_id']]['students'].append(dict(row))
            # 内层通过 (class_id,id) 索引定位每班末尾，外层只读取入选记录。
            for row in db.execute('''SELECT c.id AS class_id,r.time,r.results
                    FROM classroom c JOIN draw_record r ON r.id IN (
                        SELECT id FROM draw_record WHERE class_id=c.id ORDER BY id DESC LIMIT 50)
                    ORDER BY c.id,r.id DESC'''):
                by_id[row['class_id']]['history'].append(
                    dict(time=row['time'], results=json.loads(row['results'])))
            last = db.execute("SELECT value FROM app_setting WHERE key='last_class'").fetchone()
            return dict(classes=classes, settings=self.config(db), defaults=DEFAULTS,
                        last_class=json.loads(last[0]) if last else None)

    def history_page(self, cid, page=1):
        """按班级分页读取全部保留记录，每页 20 次，最新在前；不受首页 50 条快照限制。

        cid 为现存班级整数 ID，page 为从 1 开始的正整数；无记录返回空列表。
        时间与姓名/学号来自抽取时快照，删除或编辑学生不改写历史。
        """
        if type(cid) is not int or type(page) is not int or page < 1:
            raise ValueError('班级或页码无效')
        with self.lock, self.connect() as db:
            # 页码总数和当前页必须属于同一快照，避免独立实例清理历史时页码跳变。
            db.execute('BEGIN')
            if not db.execute('SELECT 1 FROM classroom WHERE id=?', (cid,)).fetchone():
                raise ValueError('班级不存在')
            total = db.execute('SELECT COUNT(*) FROM draw_record WHERE class_id=?', (cid,)).fetchone()[0]
            page = min(page, max(1, (total + 19) // 20))
            rows = [dict(row) for row in db.execute(
                'SELECT id,time,results FROM draw_record WHERE class_id=? ORDER BY id DESC LIMIT 20 OFFSET ?',
                (cid, (page - 1) * 20),
            )]
            for row in rows:
                row['results'] = json.loads(row['results'])
            return dict(records=rows, total=total, page=page)

    def backup(self, destination=None, kind='manual'):
        """Backup API 创建一致性副本，临时文件完成校验后替换目标。"""
        with self.lock:
            destination = Path(destination) if destination else self.backups / f'{kind}-{datetime.now():%Y%m%d-%H%M%S-%f}.db'
            if destination.resolve() == self.path.resolve():
                raise ValueError('不能覆盖正在使用的数据库')
            temp = destination.with_suffix(destination.suffix + '.tmp')
            try:
                with self.connect() as source:
                    target = sqlite3.connect(temp)
                    try:
                        source.backup(target)
                        if target.execute('PRAGMA integrity_check').fetchone()[0] != 'ok':
                            raise ValueError('备份校验失败')
                    finally:
                        target.close()
                temp.replace(destination)
            finally:
                temp.unlink(missing_ok=True)
            if kind in ('daily', 'protect'):
                with self.connect() as db:
                    count = self.config(db)['backup_count'] if kind == 'daily' else 3
                for old in sorted(self.backups.glob(f'{kind}-*.db'), reverse=True)[count:]:
                    try:
                        old.unlink()
                    except PermissionError:
                        # Windows 恢复时源备份仍被只读连接占用；保留到下次清理，
                        # 不能因删除旧文件失败而阻断已经成功的备份或恢复流程。
                        pass
            return str(destination)

    def restore(self, source):
        """恢复 v1/v2 备份；校验后在临时副本升级排序结构，保护备份后原子替换。"""
        with self.lock:
            source = Path(source).resolve()
            if source == self.path.resolve():
                raise ValueError('请选择备份文件，不能恢复正在使用的数据库')
            # SQLite 连接的事务上下文不会关闭句柄；显式 closing 避免 Windows 文件锁残留。
            with closing(sqlite3.connect(f'{source.as_uri()}?mode=ro', uri=True)) as db:
                # 校验和复制同一版本；其他程序修改备份文件时不能校验旧数据、复制新数据。
                db.execute('BEGIN')
                if db.execute('PRAGMA integrity_check').fetchone()[0] != 'ok' or db.execute('PRAGMA user_version').fetchone()[0] not in (1, 2):
                    raise ValueError('文件损坏或不是兼容的备份')
                # 只接受应用数据表；未知视图/触发器会在恢复后改变业务行为。
                allowed_tables = {'classroom', 'student', 'draw_record', 'app_setting'}
                schema = db.execute("SELECT type,name FROM sqlite_master WHERE name NOT LIKE 'sqlite_%'").fetchall()
                if any((kind == 'table' and name not in allowed_tables)
                       or kind == 'view' or (kind == 'trigger' and name != 'student_append_position')
                       for kind, name in schema):
                    raise ValueError('备份包含不支持的数据库结构')
                for table, columns in {'classroom':'id,name,position,rules,latest,request', 'student':'id,class_id,name,no,enabled,drawn,excluded', 'draw_record':'id,class_id,request,time,results', 'app_setting':'key,value'}.items():
                    db.execute(f'SELECT {columns} FROM {table} LIMIT 1')
                if db.execute('PRAGMA foreign_key_check').fetchone():
                    raise ValueError('备份关联数据不完整')
                row = db.execute("SELECT value FROM app_setting WHERE key='config'").fetchone()
                valid_settings(json.loads(row[0]) if row else {})
                last = db.execute("SELECT value FROM app_setting WHERE key='last_class'").fetchone()
                if last:
                    value = json.loads(last[0])
                    if value is not None and type(value) is not int:
                        raise ValueError('备份当前班级标识无效')
                for name, rules, latest in db.execute('SELECT name,rules,latest FROM classroom'):
                    if not isinstance(name, str) or not name.strip() or len(name) > 60:
                        raise ValueError('备份班级名称无效')
                    valid_settings(json.loads(rules), True)
                    self.validate_result_snapshot(latest)
                # integrity_check 只检查物理结构；逐项检查应用会直接渲染的历史/名单内容。
                for timestamp, results in db.execute('SELECT time,results FROM draw_record'):
                    parsed = datetime.fromisoformat(timestamp)
                    if parsed.tzinfo is None:
                        raise ValueError('备份记录时间缺少时区')
                    self.validate_result_snapshot(results)
                for name, no, enabled, drawn, excluded in db.execute('SELECT name,no,enabled,drawn,excluded FROM student'):
                    self.student_fields({'name': name, 'no': no})
                    if any(value not in (0, 1) for value in (enabled, drawn, excluded)):
                        raise ValueError('备份学生状态无效')
                self.backup(kind='protect')
                temp = self.path.with_suffix('.restore')
                try:
                    target = sqlite3.connect(temp)
                    try:
                        db.backup(target)
                        with target:
                            self.migrate_student_order(target)
                    finally:
                        target.close()
                    temp.replace(self.path)
                finally:
                    temp.unlink(missing_ok=True)

    @staticmethod
    def validate_result_snapshot(encoded):
        """拒绝无法渲染的抽取快照；兼容已删除学生的历史 ID，不要求关联现存名单。

        Args:
            encoded: 数据库中的 JSON 字符串，每次最多 100 人，空列表合法。
        Raises:
            ValueError: JSON、ID、姓名或学号类型不符合历史快照协议。
        """
        results = json.loads(encoded)
        if not isinstance(results, list) or len(results) > 100:
            raise ValueError('备份结果数据无效')
        for student in results:
            if not isinstance(student, dict) or type(student.get('id')) is not int:
                raise ValueError('备份结果学生标识无效')
            Store.student_fields(student)

    def mutate(self, action, data):
        """白名单业务写操作；SQLite 唯一约束作最终兜底，错误回滚整次变更。"""
        if not isinstance(data, dict):
            raise ValueError('操作参数须为对象')
        actions = {'save_class', 'delete_class', 'move_class', 'save_student', 'add_students',
                   'move_student', 'delete_student', 'delete_students', 'exclusions', 'reset',
                   'settings', 'rules', 'last_class', 'clear_history', 'import'}
        if not isinstance(action, str) or action not in actions:
            raise ValueError('不支持的操作')
        cid = data.get('class_id')
        if (action not in ('save_class', 'settings', 'clear_history') or cid is not None):
            if type(cid) is not int or cid <= 0:
                raise ValueError('请选择有效班级')
        for key in ('replace', 'restore_excluded'):
            if key in data and type(data[key]) is not bool:
                raise ValueError('操作开关须为布尔值')
        if action in ('save_student', 'delete_student', 'move_student'):
            sid = data.get('id')
            if sid is not None or action != 'save_student':
                if type(sid) is not int or sid <= 0:
                    raise ValueError('请选择有效学生')
        with self.lock:
            with self.connect() as check:
                cfg = self.config(check)
            if cfg['auto_backup'] and not list(self.backups.glob(f'daily-{datetime.now():%Y%m%d}-*.db')):
                self.backup(kind='daily')
            if action in ('delete_class', 'delete_student', 'delete_students', 'import', 'clear_history'):
                self.backup(kind='protect')
            try:
                with self.connect() as db:
                    # 所有先读后写的排序/校验均先取得 SQLite 写锁，不能仅靠实例 RLock。
                    db.execute('BEGIN IMMEDIATE')
                    cid = data.get('class_id')
                    if cid is not None and not db.execute('SELECT id FROM classroom WHERE id=?', (cid,)).fetchone():
                        raise ValueError('班级不存在，请刷新页面')
                    if action == 'save_class':
                        name = data.get('name')
                        if not isinstance(name, str):
                            raise ValueError('班级名称须为文本')
                        name = name.strip()
                        if not name or len(name) > 60:
                            raise ValueError('班级名称须为 1—60 个字符')
                        if cid:
                            db.execute('UPDATE classroom SET name=? WHERE id=?', (name,cid))
                        else:
                            db.execute('INSERT INTO classroom(name,position) VALUES(?,(SELECT COALESCE(MAX(position),0)+1 FROM classroom))', (name,))
                    elif action == 'delete_class':
                        db.execute('DELETE FROM classroom WHERE id=?', (cid,))
                    elif action == 'move_class':
                        direction = data.get('direction')
                        if type(direction) is not int or direction not in (-1,1):
                            raise ValueError('排序方向无效')
                        ids = [r[0] for r in db.execute('SELECT id FROM classroom ORDER BY position,id')]
                        i=ids.index(cid); j=i+direction
                        if 0 <= j < len(ids):
                            ids[i],ids[j]=ids[j],ids[i]
                            db.executemany('UPDATE classroom SET position=? WHERE id=?', enumerate(ids))
                    elif action == 'move_student':
                        # 每次只与本班完整名单的相邻学生交换，筛选时前端禁用排序。
                        # 在同一锁和事务内读取/交换，边界操作无副作用，不改抽取数据。
                        direction, sid = data.get('direction'), data.get('id')
                        if type(direction) is not int or direction not in (-1, 1):
                            raise ValueError('排序方向无效')
                        if type(cid) is not int or type(sid) is not int:
                            raise ValueError('请选择班级和学生')
                        ids = [row[0] for row in db.execute(
                            'SELECT id FROM student WHERE class_id=? ORDER BY position,id', (cid,)
                        )]
                        if sid not in ids:
                            raise ValueError('学生不存在')
                        index = ids.index(sid)
                        other = index + direction
                        if 0 <= other < len(ids):
                            ids[index], ids[other] = ids[other], ids[index]
                            db.executemany('UPDATE student SET position=? WHERE id=? AND class_id=?',
                                           [(i, student_id, cid) for i, student_id in enumerate(ids, 1)])
                    elif action == 'save_student':
                        name, no = self.student_fields(data)
                        enabled = data.get('enabled', True)
                        if type(enabled) is not bool:
                            raise ValueError('参与状态无效')
                        sid = data.get('id')
                        if sid:
                            cursor=db.execute('UPDATE student SET name=?,no=?,enabled=? WHERE id=? AND class_id=?', (name,no,int(enabled),sid,cid))
                            if not cursor.rowcount: raise ValueError('学生不存在')
                        else:
                            db.execute('INSERT INTO student(class_id,name,no,enabled) VALUES(?,?,?,?)', (cid,name,no,int(enabled)))
                    elif action == 'add_students':
                        # 表格新增按整批提交，任一行失败都会回滚整批。
                        # 学号留空存 NULL，可多行为空；只有手填的非空学号参与班内唯一校验。
                        rows = data.get('rows')
                        if type(cid) is not int or cid <= 0:
                            raise ValueError('请选择班级')
                        if not isinstance(rows, list) or not 1 <= len(rows) <= 100:
                            raise ValueError('每次新增 1—100 名学生')
                        used = {row[0] for row in db.execute(
                            'SELECT no FROM student WHERE class_id=? AND no IS NOT NULL', (cid,)
                        )}
                        clean = []
                        for index, row in enumerate(rows, 1):
                            if not isinstance(row, dict):
                                raise ValueError(f'第 {index} 行学生信息格式无效')
                            try:
                                name, no = self.student_fields(row)
                            except ValueError as exc:
                                raise ValueError(f'第 {index} 行：{exc}') from exc
                            enabled = row.get('enabled', True)
                            if type(enabled) is not bool:
                                raise ValueError(f'第 {index} 行参与状态无效')
                            if no is not None:
                                if no in used:
                                    raise ValueError(f'第 {index} 行：学号 {no} 重复')
                                used.add(no)
                            clean.append((name, no, enabled))
                        db.executemany(
                            'INSERT INTO student(class_id,name,no,enabled) VALUES(?,?,?,?)',
                            [(cid, name, no, int(enabled)) for name, no, enabled in clean],
                        )
                    elif action == 'delete_student':
                        db.execute('DELETE FROM student WHERE id=? AND class_id=?', (data.get('id'),cid))
                    elif action == 'delete_students':
                        ids = data.get('ids', [])
                        valid = {r[0] for r in db.execute('SELECT id FROM student WHERE class_id=?', (cid,))}
                        if not isinstance(ids, list) or any(type(i) is not int or i not in valid for i in ids):
                            raise ValueError('删除名单无效')
                        db.executemany('DELETE FROM student WHERE id=? AND class_id=?', [(i, cid) for i in ids])
                    elif action == 'exclusions':
                        ids = data.get('ids', [])
                        valid = {r[0] for r in db.execute('SELECT id FROM student WHERE class_id=?', (cid,))}
                        if not isinstance(ids,list) or any(type(i) is not int or i not in valid for i in ids):
                            raise ValueError('排除名单无效')
                        db.execute('UPDATE student SET excluded=0 WHERE class_id=?', (cid,))
                        db.executemany('UPDATE student SET excluded=1 WHERE id=? AND class_id=?', [(i,cid) for i in ids])
                    elif action == 'reset':
                        db.execute('UPDATE student SET drawn=0 WHERE class_id=?', (cid,))
                        if data.get('restore_excluded'): db.execute('UPDATE student SET excluded=0 WHERE class_id=?', (cid,))
                        db.execute("UPDATE classroom SET latest='[]',request=NULL WHERE id=?", (cid,))
                    elif action == 'settings':
                        config = valid_settings(data.get('values'))
                        db.execute("INSERT OR REPLACE INTO app_setting VALUES('config',?)", (json.dumps(config),))
                        if config['history_days']:
                            cutoff=(datetime.now(timezone.utc)-timedelta(days=config['history_days'])).isoformat()
                            db.execute('DELETE FROM draw_record WHERE time<?', (cutoff,))
                    elif action == 'rules':
                        rules = valid_settings(data.get('values'), True)
                        db.execute('UPDATE classroom SET rules=? WHERE id=?', (json.dumps(rules),cid))
                    elif action == 'last_class':
                        db.execute("INSERT OR REPLACE INTO app_setting VALUES('last_class',?)", (json.dumps(cid),))
                    elif action == 'clear_history':
                        db.execute('DELETE FROM draw_record')
                    elif action == 'import':
                        rows=data.get('rows', [])
                        if not isinstance(rows,list) or not rows or len(rows)>10000:
                            raise ValueError('每次导入 1—10000 行')
                        clean=[self.student_fields(r) for r in rows]
                        numbers=[no for _,no in clean if no]
                        if len(numbers)!=len(set(numbers)):
                            raise ValueError('导入学号重复，请修改后重试')
                        if data.get('replace'):
                            db.execute('DELETE FROM student WHERE class_id=?', (cid,))
                            db.execute("UPDATE classroom SET latest='[]',request=NULL WHERE id=?", (cid,))
                        # 空学号保持 NULL；沿用唯一约束校验非空学号，失败时替换操作也会回滚。
                        db.executemany('INSERT INTO student(class_id,name,no) VALUES(?,?,?)', [(cid,*r) for r in clean])
                    else:
                        raise ValueError('不支持的操作')
            except sqlite3.IntegrityError as exc:
                raise ValueError('班级名称或本班学生学号重复，请检查') from exc
            return self.snapshot()

    @staticmethod
    def student_fields(data):
        """姓名允许重名，学号保留前导零；缺省、空白学号统一为 None，不自动生成。"""
        if not isinstance(data, dict):
            raise ValueError('学生信息须为对象')
        # null 姓名不能被 str 转成有效姓名 "None"；数值 0 学号也不能当空值丢掉。
        name = data.get('name')
        number = data.get('no')
        if not isinstance(name, str) or isinstance(number, (bool, list, dict)):
            raise ValueError('姓名须为文本，学号须为文本或数字')
        name = name.strip()
        no = str(number).strip() if number is not None else None
        no = no or None
        if not name or len(name)>60 or (no and len(no)>60):
            raise ValueError('姓名不能为空，姓名和学号最多 60 个字符')
        return name,no

    def draw(self, cid, request):
        """系统随机源无放回抽样；请求 ID 相同返回最近结果，避免即时重试重复抽取。"""
        if type(cid) is not int or cid <= 0:
            raise ValueError('请选择有效班级')
        with self.lock:
            with self.connect() as db:
                row=db.execute('SELECT * FROM classroom WHERE id=?', (cid,)).fetchone()
                if not row: raise ValueError('请先创建班级')
                if not isinstance(request,str) or not 8<=len(request)<=100: raise ValueError('请求标识无效')
                if row['request']==request: return self.snapshot()
                cfg={**self.config(db),**json.loads(row['rules'])}
            if cfg['auto_backup'] and not list(self.backups.glob(f'daily-{datetime.now():%Y%m%d}-*.db')):
                self.backup(kind='daily')
            with self.connect() as db:
                # 独立 EXE/Store 的 Python 锁不共享，必须先取得 SQLite 写锁再读候选。
                # 同一请求在锁内再次校验，避免并发窗口重复提交或无放回模式抽中同一人。
                db.execute('BEGIN IMMEDIATE')
                row = db.execute('SELECT request,rules FROM classroom WHERE id=?', (cid,)).fetchone()
                if not row:
                    raise ValueError('班级不存在')
                if row['request'] == request:
                    return self.snapshot()
                cfg = {**self.config(db), **json.loads(row['rules'])}
                pool=[dict(s) for s in db.execute('SELECT * FROM student WHERE class_id=? AND enabled=1 AND excluded=0', (cid,)) if cfg['mode']=='random' or not s['drawn']]
                if len(pool)<cfg['count']: raise ValueError(f"当前可抽 {len(pool)} 人，不足 {cfg['count']} 人。请重置本轮或调整人数。")
                selected=secrets.SystemRandom().sample(pool,cfg['count'])
                # 抽取快照保留真实学号；内部主键仅作记录标识，不得替代未填写的学号。
                results=[dict(id=s['id'],name=s['name'],no=s['no']) for s in selected]
                if cfg['mode']=='unique':
                    db.executemany('UPDATE student SET drawn=1 WHERE id=?', [(s['id'],) for s in selected])
                encoded=json.dumps(results,ensure_ascii=False)
                db.execute('UPDATE classroom SET latest=?,request=? WHERE id=?',(encoded,request,cid))
                if cfg['save_history']:
                    db.execute('INSERT INTO draw_record(class_id,request,time,results) VALUES(?,?,?,?)',(cid,request,datetime.now(timezone.utc).isoformat(),encoded))
                if cfg['history_days']:
                    db.execute('DELETE FROM draw_record WHERE time<?',((datetime.now(timezone.utc)-timedelta(days=cfg['history_days'])).isoformat(),))
            return self.snapshot()

    @staticmethod
    def read_file(path):
        """读取预览二维数组；CSV 回退 GB18030，确认导入由另一步事务提交。

        原文件最多 10 MiB，XLSX 解压声明总大小最多 64 MiB；逐行限制每表
        10001 行（含一行表头）、50 列，不信任可缺失或偏小的 XLSX 尺寸元数据。
        XLSX 公式拒绝导入；XLS 由 xlrd 读取缓存值，不执行公式。
        """
        path=Path(path)
        if path.stat().st_size>10*1024*1024: raise ValueError('文件请小于 10 MB')
        if path.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(path) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 64 * 1024 * 1024:
                    raise ValueError('工作簿解压后内容过大，请拆分后导入')
            book=load_workbook(path,read_only=True,data_only=False)
            try:
                sheets={}
                for sheet in book:
                    if (sheet.max_row or 0)>10001 or (sheet.max_column or 0)>50:
                        raise ValueError('工作表最多 10000 行、50 列')
                    # 部分合法生成器不写 dimension；偏小的 dimension 也不能截断真实数据。
                    sheet.reset_dimensions()
                    rows=[]
                    for row_number, cells in enumerate(sheet.iter_rows(), 1):
                        if row_number > 10001 or len(cells) > 50:
                            raise ValueError('工作表最多 10000 行、50 列')
                        if any(c.data_type=='f' for c in cells): raise ValueError('包含公式，请转为普通值后导入')
                        rows.append(['' if c.value is None else str(c.value) for c in cells])
                    sheets[sheet.title]=rows
                return sheets
            finally: book.close()
        if path.suffix.lower()=='.xls':
            import xlrd
            book=xlrd.open_workbook(path, on_demand=True)
            try:
                sheets={}
                for sheet in book.sheets():
                    if sheet.nrows>10001 or sheet.ncols>50: raise ValueError('工作表最多 10000 行、50 列')
                    sheets[sheet.name]=[
                        [str(int(cell.value)) if cell.ctype == xlrd.XL_CELL_NUMBER and cell.value.is_integer()
                         else str(cell.value) if cell.value is not None else '' for cell in sheet.row(row)]
                        for row in range(sheet.nrows)
                    ]
                return sheets
            finally:
                book.release_resources()
        if path.suffix.lower()!='.csv': raise ValueError('仅支持 .xlsx、.xls 或 .csv')
        try: content=path.read_text(encoding='utf-8-sig')
        except UnicodeDecodeError: content=path.read_text(encoding='gb18030')
        rows = []
        for row in csv.reader(io.StringIO(content)):
            if len(rows) >= 10001: raise ValueError('最多导入 10000 行')
            if len(row) > 50: raise ValueError('工作表最多 50 列')
            rows.append(row)
        return {'CSV': rows}
