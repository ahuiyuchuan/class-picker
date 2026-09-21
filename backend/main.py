"""Windows 桌面入口与窗口桥接。

@author ahui
窗口状态以 WinForms 实际状态为准，避免最大化、最小化、全屏互相污染。
"""
from pathlib import Path
import functools
import http.server
import socketserver
import sys
import threading
import os
import logging
from storage import Store


class WindowApi:
    """暴露窗口操作及白名单本地数据服务；窗口状态在 GUI 线程读写。"""

    def __init__(self):
        self._window = None
        self._before_fullscreen = False
        self._normal_before_fullscreen = None
        self._changing_fullscreen = False
        folder = os.environ.get("CLASS_PICKER_DATA_DIR") or str(Path(os.environ["LOCALAPPDATA"]) / "ClassroomPicker" / "data")
        self._store = Store(folder)
        logging.basicConfig(filename=str(self._store.folder / 'errors.log'),
                            encoding='utf-8', level=logging.ERROR)

    def data(self, action, payload=None):
        """白名单数据桥接；可预期校验错误显示给用户，未知异常写本地诊断日志。"""
        try:
            payload = {} if payload is None else payload
            if not isinstance(action, str) or not isinstance(payload, dict):
                raise ValueError('操作名称须为文本，参数须为对象')
            if action == 'load':
                result = self._store.snapshot()
            elif action == 'history_page':
                result = self._store.history_page(payload.get('class_id'), payload.get('page', 1))
            elif action == 'draw':
                result = self._store.draw(payload.get('class_id'), payload.get('request'))
            elif action == 'choose_import':
                import webview
                files = self._window.create_file_dialog(
                    webview.FileDialog.OPEN,
                    allow_multiple=False,
                    file_types=('Excel 工作簿 (*.xlsx)', 'Excel97 工作簿 (*.xls)', 'CSV 文件 (*.csv)'),
                )
                result = self._store.read_file(files[0]) if files else None
            elif action == 'download_template':
                import webview
                path = self._window.create_file_dialog(
                    webview.FileDialog.SAVE,
                    save_filename='班级名单导入模板.xlsx',
                    file_types=('Excel 工作簿 (*.xlsx)',),
                )
                if path:
                    path = path[0] if isinstance(path, (tuple, list)) else path
                    if Path(path).suffix.lower() != '.xlsx':
                        path = f'{path}.xlsx'
                    from openpyxl import Workbook
                    book = Workbook()
                    sheet = book.active
                    sheet.title = '学生名单'
                    sheet.append(['姓名', '学号'])
                    # 学号以文本写入，Excel 不会改写为数值或科学计数法。
                    for row in [('张三', '202609201940'), ('李四', '202609201941'), ('王五', '202609201942')]:
                        sheet.append(row)
                    book.save(path)
                result = bool(path)
            elif action == 'export':
                import webview
                classes = self._store.snapshot()['classes']
                selected = next((c for c in classes if c['id'] == payload.get('class_id')), None)
                if not selected: raise ValueError('请选择班级')
                path = self._window.create_file_dialog(
                    webview.FileDialog.SAVE,
                    save_filename='学生名单.xlsx',
                    file_types=('Excel 工作簿 (*.xlsx)',),
                )
                if path:
                    path = path[0] if isinstance(path, (tuple, list)) else path
                    if Path(path).suffix.lower() != '.xlsx':
                        path = f'{path}.xlsx'
                    rows = [['姓名', '学号']] + [[s['name'], s['no']] for s in selected['students']]
                    from openpyxl import Workbook
                    book = Workbook()
                    sheet = book.active
                    sheet.title = '学生名单'
                    for row_index, row in enumerate(rows, 1):
                        # XLSX 明确标记字符串即可阻止公式解释，不在用户数据前插入单引号。
                        # 用已知行列直接定位，避免逐行计算 max_row/max_column 扫描整张表。
                        for column_index, value in enumerate(row, 1):
                            cell = sheet.cell(row_index, column_index, value)
                            if cell.value is not None:
                                cell.data_type = 's'
                                cell.number_format = '@'
                    book.save(path)
                result = bool(path)
            elif action == 'backup':
                import webview
                path = self._window.create_file_dialog(webview.FileDialog.SAVE, save_filename='班级抽人备份.db', file_types=('数据备份 (*.db)',))
                if path:
                    self._store.backup(path[0] if isinstance(path, (tuple, list)) else path)
                result = bool(path)
            elif action == 'restore':
                import webview
                files = self._window.create_file_dialog(webview.FileDialog.OPEN, allow_multiple=False, file_types=('数据备份 (*.db)',))
                if files: self._store.restore(files[0])
                result = self._store.snapshot() if files else None
            elif action == 'open_data':
                os.startfile(str(self._store.folder))
                result = True
            else:
                result = self._store.mutate(action, payload)
            return {'ok': True, 'value': result}
        except (ValueError, OSError) as exc:
            return {'ok': False, 'error': str(exc)}
        except Exception:
            logging.exception('Local data operation failed: %s', action)
            return {'ok': False, 'error': '操作失败，数据未确认保存。请重试或检查本地数据目录权限。'}

    def state(self):
        """返回真实窗口状态，供按钮图标和全屏布局同步。"""
        from System import Func, Object
        result = {}

        def read():
            form = self._window.native
            result.update(maximized=str(form.WindowState) == "Maximized",
                          minimized=str(form.WindowState) == "Minimized",
                          fullscreen=bool(form.is_fullscreen))
            return None

        self._window.native.Invoke(Func[Object](read))
        return result

    def minimize(self):
        """无论普通或最大化状态都最小化；不额外 restore，保留系统还原语义。"""
        self._window.minimize()
        return self.state()

    def _set_maximized_bounds(self):
        """仅在 GUI 线程调用；按所在显示器工作区限制无边框窗口最大化。

        WinForms 将该矩形传给 WM_GETMINMAXINFO，位置必须相对显示器原点，
        尺寸使用原生物理像素，不能混入前端 CSS 像素或再次应用 DPI 缩放。
        全屏有意使用整个显示器，因此由 fullscreen 暂时解除该约束。
        """
        from System.Drawing import Rectangle
        from System.Windows.Forms import Screen
        form = self._window.native
        screen = Screen.FromControl(form)
        work = screen.WorkingArea
        form.MaximizedBounds = Rectangle(
            work.X - screen.Bounds.X, work.Y - screen.Bounds.Y, work.Width, work.Height)

    def _initialize_window(self):
        """显示后安装工作区约束；拖到另一显示器后刷新，系统最大化也适用。"""
        from System import Func, Object
        from System.Windows.Forms import FormWindowState

        def moved(sender, event):
            form = self._window.native
            if (not self._changing_fullscreen and not form.is_fullscreen
                    and form.WindowState == FormWindowState.Normal):
                self._set_maximized_bounds()

        def initialize():
            self._set_maximized_bounds()
            self._window.native.LocationChanged += moved
            return None

        self._window.native.Invoke(Func[Object](initialize))

    def maximize(self):
        """在 GUI 线程原子切换最大化/还原，不依赖本地布尔推测。"""
        from System import Func, Object
        from System.Windows.Forms import FormWindowState

        def toggle():
            form = self._window.native
            if not form.is_fullscreen:
                self._set_maximized_bounds()
                form.WindowState = (FormWindowState.Normal
                                    if form.WindowState == FormWindowState.Maximized
                                    else FormWindowState.Maximized)
            return None

        self._window.native.Invoke(Func[Object](toggle))
        return self.state()

    def fullscreen(self):
        """原子切换全屏，退出时恢复原状态以及真正的普通窗口位置和尺寸。

        pywebview 只保存进入全屏时的当前尺寸；从最大化进入时，该尺寸不是
        普通窗口的还原尺寸。必须独立保存 RestoreBounds，并在退出后先恢复
        普通矩形再最大化，否则下一次点击还原只改状态、不会缩小窗口。
        过渡期间合并 Dock 布局并暂停重绘，WebView 只接收最终尺寸，避免缩放闪动。
        返回同步后的原生状态；不改变数据库或用户配置。
        """
        from System import Func, Object
        from System.Drawing import Rectangle
        from System.Windows.Forms import FormWindowState, Screen
        from window_transition import batch_window_transition

        def toggle():
            form = self._window.native
            if not form.is_fullscreen:
                self._before_fullscreen = form.WindowState == FormWindowState.Maximized
                self._normal_before_fullscreen = (
                    form.RestoreBounds if self._before_fullscreen else form.Bounds)
                form.MaximizedBounds = Rectangle.Empty
                form.toggle_fullscreen()
            else:
                form.toggle_fullscreen()
                form.WindowState = FormWindowState.Normal
                if self._normal_before_fullscreen is not None:
                    normal = self._normal_before_fullscreen
                    # 全屏期间拔掉显示器时，保留库的可见性保证，不能把窗口还原到屏幕外。
                    if not any(screen.WorkingArea.IntersectsWith(normal) for screen in Screen.AllScreens):
                        work = Screen.FromControl(form).WorkingArea
                        normal = Rectangle(work.X, work.Y,
                                           min(normal.Width, work.Width), min(normal.Height, work.Height))
                    form.Bounds = normal
                self._set_maximized_bounds()
                if self._before_fullscreen:
                    form.WindowState = FormWindowState.Maximized
                self._normal_before_fullscreen = None
            return None

        # LocationChanged 会在库的全屏过渡中同步触发；此时不得重新施加工作区限制。
        def transition():
            self._changing_fullscreen = True
            try:
                with batch_window_transition(self._window.native):
                    return toggle()
            finally:
                self._changing_fullscreen = False

        self._window.native.Invoke(Func[Object](transition))
        return self.state()

    def close(self):
        """关闭窗口，由入口 finally 释放静态服务。"""
        self._window.destroy()


class StaticHandler(http.server.SimpleHTTPRequestHandler):
    """单文件无控制台模式 sys.stderr 为 None，禁止默认日志写入导致请求失败。"""

    def log_message(self, format, *args):
        pass


def resource_root():
    """返回源码根目录或单文件解包目录。"""
    return Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parents[1]))


def main():
    """只监听本机的资源服务与 WebView2 窗口共用生命周期。"""
    import webview
    root = resource_root()
    handler = functools.partial(StaticHandler, directory=str(root / 'frontend' / 'dist'))
    server = socketserver.ThreadingTCPServer(('127.0.0.1', 0), handler)
    server.daemon_threads = True
    threading.Thread(target=server.serve_forever, daemon=True, name='static-assets').start()
    api = WindowApi()
    api._window = webview.create_window(
        '班级随机抽人', f'http://127.0.0.1:{server.server_address[1]}/index.html',
        width=1280, height=820, min_size=(800, 560), text_select=False,
        frameless=True, easy_drag=False, js_api=api,
    )
    # Only the separate brand region receives drag events; system buttons are siblings.
    api._window.events.shown += api._initialize_window
    webview.settings['DRAG_REGION_SELECTOR'] = '.window-drag-region'
    try:
        if '--verify-window' in sys.argv:
            from verify_window import run
            webview.start(run, (api,), gui='edgechromium', icon=str(root / 'packaging/ClassPicker.ico'))
        else:
            webview.start(gui='edgechromium', icon=str(root / 'packaging/ClassPicker.ico'))
    finally:
        server.shutdown()
        server.server_close()


if __name__ == '__main__':
    main()
