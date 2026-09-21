"""@author ahui — 隔离临时数据库回归，覆盖真实事务与重启，不读写用户名单。"""
import sys
import tempfile
import unittest
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from storage import Store, DEFAULTS

class StoreTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.store=Store(Path(self.temp.name)/'data')
        snap=self.store.mutate('save_class',{'name':'测试一班'})
        self.cid=snap['classes'][0]['id']
        self.store.mutate('import',{'class_id':self.cid,'rows':[{'name':f'学生{i}','no':f'{i:02}'} for i in range(1,11)]})

    def tearDown(self):
        self.temp.cleanup()

    def test_restart_and_request_idempotency(self):
        result=self.store.draw(self.cid,'request-0001')
        again=self.store.draw(self.cid,'request-0001')
        self.assertEqual(result,again)
        restarted=Store(self.store.folder).snapshot()
        self.assertEqual(result,restarted)
        self.assertEqual(len(result['classes'][0]['history']),1)

    def test_atomic_import_and_preserved_numbers(self):
        with self.assertRaises(ValueError):
            self.store.mutate('import',{'class_id':self.cid,'rows':[{'name':'new','no':'11'},{'name':'duplicate','no':'01'}]})
        students=self.store.snapshot()['classes'][0]['students']
        self.assertEqual(len(students),10)
        self.assertEqual(students[0]['no'],'01')

    def test_config_inheritance_and_count(self):
        self.store.mutate('rules',{'class_id':self.cid,'values':{'count':2}})
        self.store.mutate('settings',{'values':{**DEFAULTS,'count':3}})
        result=self.store.draw(self.cid,'request-0002')['classes'][0]['latest']
        self.assertEqual(len(result),2)
        self.assertEqual(len({x['id'] for x in result}),2)

    def test_concurrent_same_request(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            results=list(pool.map(lambda _:self.store.draw(self.cid,'same-request'),range(4)))
        self.assertTrue(all(results[0]==r for r in results))
        self.assertEqual(len(self.store.snapshot()['classes'][0]['history']),1)

    def test_excluded_disabled_and_exhaustion(self):
        students=self.store.snapshot()['classes'][0]['students']
        self.store.mutate('exclusions',{'class_id':self.cid,'ids':[s['id'] for s in students[:8]]})
        self.store.mutate('save_student',{**students[8],'enabled':False})
        result=self.store.draw(self.cid,'last-person')
        self.assertEqual(result['classes'][0]['latest'][0]['id'],students[9]['id'])
        with self.assertRaises(ValueError):self.store.draw(self.cid,'no-more-people')
        self.assertEqual(len(self.store.snapshot()['classes'][0]['history']),1)

    def test_backup_restore_and_invalid_backup(self):
        path=self.store.backup(Path(self.temp.name)/'backup.db')
        self.store.mutate('delete_class',{'class_id':self.cid})
        self.store.restore(path)
        self.assertEqual(len(self.store.snapshot()['classes'][0]['students']),10)
        self.assertTrue(list(self.store.backups.glob('protect-*.db')))

    def test_random_and_no_history(self):
        self.store.mutate('settings',{'values':{**DEFAULTS,'mode':'random','save_history':False}})
        snap=self.store.draw(self.cid,'random-request')['classes'][0]
        self.assertEqual(snap['history'],[])
        self.assertFalse(any(s['drawn'] for s in snap['students']))
        self.assertEqual(len(snap['latest']),1)

    def test_reset_and_class_isolation(self):
        self.store.mutate('save_class',{'name':'二班'})
        self.store.draw(self.cid,'reset-request')
        self.store.mutate('reset',{'class_id':self.cid,'restore_excluded':True})
        snap=self.store.snapshot()
        self.assertEqual(snap['classes'][0]['latest'],[])
        self.assertEqual(len(snap['classes'][0]['history']),1)
        self.assertEqual(snap['classes'][1]['latest'],[])

    def test_add_students_batch_numbers_and_flags(self):
        """空学号保持 NULL，手填学号的前导零和参与状态保持原值。"""
        snapshot = self.store.mutate('add_students', {'class_id': self.cid, 'rows': [
            {'name': '空学号', 'no': ''},
            {'name': '手填学号', 'no': '1', 'enabled': False},
            {'name': '文本学号', 'no': '0007'},
        ]})
        students = snapshot['classes'][0]['students']
        self.assertEqual(len(students), 13)
        by_name = {student['name']: student for student in students}
        self.assertIsNone(by_name['空学号']['no'])
        self.assertEqual(by_name['手填学号']['enabled'], 0)
        self.assertEqual(by_name['文本学号']['no'], '0007')

    def test_add_students_validation_rolls_back_entire_batch(self):
        """后续行错误不得留下前面的新增记录，也不得改动其他班级。"""
        self.store.mutate('save_class', {'name': '隔离班'})
        before = self.store.snapshot()
        cases = [
            [{'name': '先填'}, {'name': ''}],
            [{'name': '先填'}, {'name': '重复已存学号', 'no': '01'}],
            [{'name': '先填', 'no': '12'}, {'name': '批内重复', 'no': '12'}],
            [{'name': '先填'}, {'name': '状态错误', 'enabled': 'false'}],
            [],
            [{'name': '超量'}] * 101,
            ['非法行'],
        ]
        for rows in cases:
            with self.subTest(rows=rows):
                with self.assertRaises(ValueError):
                    self.store.mutate('add_students', {'class_id': self.cid, 'rows': rows})
                self.assertEqual(self.store.snapshot(), before)

    def test_add_students_concurrent_empty_numbers(self):
        """并发批次允许多个空学号，不能补号或触发唯一约束冲突。"""
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(lambda i: self.store.mutate('add_students', {
                'class_id': self.cid, 'rows': [{'name': f'批次{i}学生{j}'} for j in range(3)],
            }), range(2)))
        students = self.store.snapshot()['classes'][0]['students']
        self.assertEqual(len(students), 16)
        self.assertEqual(sum(student['no'] is None for student in students), 6)

    def test_empty_student_number_across_save_edit_import(self):
        """单条新增、编辑清空、批量导入及重启都保留空学号。"""
        self.store.mutate('save_student', {'class_id': self.cid, 'name': '未填学号'})
        students = self.store.snapshot()['classes'][0]['students']
        self.assertIsNone(students[-1]['no'])
        self.store.mutate('save_student', {
            'class_id': self.cid, 'id': students[0]['id'],
            'name': students[0]['name'], 'no': '   ',
        })
        self.store.mutate('import', {'class_id': self.cid, 'rows': [
            {'name': '导入甲', 'no': ''}, {'name': '导入乙'},
        ]})
        restarted = Store(self.store.folder).snapshot()['classes'][0]['students']
        self.assertEqual(sum(student['no'] is None for student in restarted), 4)
        self.assertEqual(restarted[1]['no'], '02')

    def test_draw_keeps_empty_student_number(self):
        """空学号参与抽取后，结果、历史和重启快照都不得用内部 ID 补号。"""
        self.store.mutate('import', {'class_id': self.cid, 'replace': True, 'rows': [
            {'name': '无学号学生'},
        ]})
        snapshot = self.store.draw(self.cid, 'empty-number-draw')['classes'][0]
        self.assertIsNone(snapshot['latest'][0]['no'])
        restarted = Store(self.store.folder).snapshot()['classes'][0]
        self.assertEqual(snapshot, restarted)
        self.assertIsNone(restarted['history'][0]['results'][0]['no'])

    def test_student_order_persistence_and_boundaries(self):
        """排序只影响本班顺序，重启/追加后保留，首尾越界和非法参数不改数据。"""
        students = self.store.snapshot()['classes'][0]['students']
        ids = [s['id'] for s in students]
        self.store.mutate('move_student', {'class_id': self.cid, 'id': ids[1], 'direction': -1})
        self.store.mutate('save_student', {'class_id': self.cid, 'name': '追加学生'})
        snapshot = Store(self.store.folder).snapshot()
        ordered = snapshot['classes'][0]['students']
        self.assertEqual([s['id'] for s in ordered[:-1]], [ids[1], ids[0], *ids[2:]])
        self.store.mutate('move_student', {'class_id': self.cid, 'id': ids[1], 'direction': -1})
        self.store.mutate('move_student', {'class_id': self.cid, 'id': ordered[-1]['id'], 'direction': 1})
        self.assertEqual(self.store.snapshot(), snapshot)
        for payload in [{'id': ids[0], 'direction': True}, {'id': 99999, 'direction': 1}]:
            with self.assertRaises(ValueError):
                self.store.mutate('move_student', {'class_id': self.cid, **payload})
        self.assertEqual(self.store.snapshot(), snapshot)
        self.store.mutate('move_student', {'class_id': self.cid, 'id': ids[1], 'direction': 1})
        self.assertEqual([s['id'] for s in self.store.snapshot()['classes'][0]['students'][:-1]], ids)

    def test_student_order_migrates_old_database_and_restore(self):
        """模拟真实 v1 结构，确认首次升级备份与恢复旧备份都补齐排序结构。"""
        with self.store.connect() as db:
            db.execute('DROP TRIGGER student_append_position')
            db.execute('DROP INDEX student_class_position')
            db.execute('ALTER TABLE student DROP COLUMN position')
            db.execute('PRAGMA user_version=1')
        legacy = self.store.backup(Path(self.temp.name) / 'legacy.db')
        migrated = Store(self.store.folder)
        self.assertTrue(list(migrated.backups.glob('pre-v2-*.db')))
        students = migrated.snapshot()['classes'][0]['students']
        self.assertEqual([s['no'] for s in students], [f'{i:02}' for i in range(1, 11)])
        migrated.mutate('move_student', {'class_id': self.cid, 'id': students[1]['id'], 'direction': -1})
        migrated.restore(legacy)
        self.assertEqual(migrated.snapshot()['classes'][0]['students'], students)
        with migrated.connect() as db:
            self.assertEqual(db.execute('PRAGMA user_version').fetchone()[0], 2)

    def test_history_pagination_and_class_isolation(self):
        """超过首页快照上限仍可翻页读取，首尾无重复，班级隔离且参数严格校验。"""
        self.store.mutate('settings', {'values': {**DEFAULTS, 'mode': 'random'}})
        for index in range(55):
            self.store.draw(self.cid, f'history-page-{index}')
        pages = [self.store.history_page(self.cid, page) for page in (1, 2, 3)]
        self.assertEqual([len(page['records']) for page in pages], [20, 20, 15])
        ids = [row['id'] for page in pages for row in page['records']]
        self.assertEqual(len(set(ids)), 55)
        self.assertEqual(ids, sorted(ids, reverse=True))
        self.assertEqual(self.store.history_page(self.cid, 999)['page'], 3)
        snap = self.store.mutate('save_class', {'name': '空记录班'})
        other = next(c['id'] for c in snap['classes'] if c['name'] == '空记录班')
        self.assertEqual(self.store.history_page(other)['total'], 0)
        for cid, page in [(self.cid, 0), (self.cid, True), (9999, 1)]:
            with self.assertRaises(ValueError):
                self.store.history_page(cid, page)

    def test_independent_stores_serialize_unique_draws(self):
        """独立 Store 的锁不共享；数据库事务必须保障并发窗口无放回抽取。"""
        stores = [Store(self.store.folder) for _ in range(4)]
        with ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(lambda i: stores[i].draw(self.cid, f'multi-window-{i}'), range(4)))
        history = self.store.snapshot()['classes'][0]['history']
        self.assertEqual(len(history), 4)
        self.assertEqual(len({row['results'][0]['id'] for row in history}), 4)
        with ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(lambda i: stores[i].draw(self.cid, 'shared-window-request'), range(4)))
        self.assertEqual(len(self.store.snapshot()['classes'][0]['history']), 5)

    def test_xls_numeric_ids_and_export_text_roundtrip(self):
        """真实 XLS 数值学号不带 .0；XLSX 特殊前缀按字符串写入，往返内容不变。"""
        import xlwt
        from types import SimpleNamespace
        from main import WindowApi
        source = Path(self.temp.name) / 'numeric.xls'
        book = xlwt.Workbook()
        sheet = book.add_sheet('名单')
        sheet.write(0, 0, '测试学生')
        sheet.write(0, 1, 202609201940)
        book.save(str(source))
        self.assertEqual(Store.read_file(source)['名单'][0][1], '202609201940')
        self.store.mutate('import', {'class_id': self.cid, 'replace': True, 'rows': [
            {'name': '=测试姓名', 'no': '+0001'}, {'name': '@测试姓名', 'no': '0'},
        ]})
        output = Path(self.temp.name) / 'roundtrip.xlsx'
        api = WindowApi.__new__(WindowApi)
        api._store = self.store
        api._window = SimpleNamespace(create_file_dialog=lambda *args, **kwargs: str(output))
        self.assertTrue(api.data('export', {'class_id': self.cid})['ok'])
        self.assertEqual(Store.read_file(output)['学生名单'][1:], [['=测试姓名', '+0001'], ['@测试姓名', '0']])

    def test_invalid_restore_keeps_current_data(self):
        """物理结构正常但历史 JSON 内容异常的备份不得替换当前名单。"""
        import sqlite3
        before = self.store.draw(self.cid, 'restore-shape-test')
        backup = Path(self.store.backup(Path(self.temp.name) / 'invalid.db'))
        db = sqlite3.connect(backup)
        try:
            with db:
                db.execute("UPDATE draw_record SET results='[null]'")
        finally:
            db.close()
        with self.assertRaises(ValueError):
            self.store.restore(backup)
        self.assertEqual(self.store.snapshot(), before)

    def test_student_validation_preserves_numeric_zero(self):
        """非法姓名不可静默转成文本；数值 0 作为真实学号保留。"""
        for row in [None, {'name': None}, {'name': '学生', 'no': []}]:
            with self.assertRaises(ValueError):
                Store.student_fields(row)
        self.assertEqual(Store.student_fields({'name': '学生', 'no': 0}), ('学生', '0'))

    def test_restore_old_protect_backup_while_pruning(self):
        """恢复最旧保护备份时，保留策略不得因 Windows 源文件占用阻止恢复。"""
        oldest = self.store.backup(kind='protect')
        self.store.backup(kind='protect')
        self.store.backup(kind='protect')
        self.store.restore(oldest)
        self.assertEqual(len(self.store.snapshot()['classes'][0]['students']), 10)

    def test_workbook_headers_use_student_number(self):
        """实际生成导入模板及导出文件，验证学号表头与文本前导零，无需系统文件对话框。"""
        from types import SimpleNamespace
        from openpyxl import load_workbook
        from main import WindowApi
        api = WindowApi.__new__(WindowApi)
        api._store = self.store
        for action in ['download_template', 'export']:
            path = Path(self.temp.name) / f'{action}.xlsx'
            api._window = SimpleNamespace(create_file_dialog=lambda *args, **kwargs: str(path))
            result = api.data(action, {'class_id': self.cid})
            self.assertTrue(result['ok'], result)
            book = load_workbook(path)
            try:
                sheet = book.active
                self.assertEqual([sheet['A1'].value, sheet['B1'].value], ['姓名', '学号'])
                if action == 'download_template':
                    self.assertEqual(list(sheet.values)[1:], [
                        ('张三', '202609201940'), ('李四', '202609201941'), ('王五', '202609201942'),
                    ])
                    self.assertEqual(sheet['B2'].data_type, 's')
                if action == 'export':
                    self.assertEqual(sheet['B2'].value, '01')
                    self.assertEqual(sheet['B2'].data_type, 's')
            finally:
                book.close()

if __name__=='__main__':unittest.main()
