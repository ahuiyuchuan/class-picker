"""数据审查回归：导入兼容性、参数边界、快照查询及跨实例排序。

@author ahui
使用独立临时数据库；并发测试在读取排序后延迟，暴露先读后加锁的丢失更新。
"""

from concurrent.futures import ThreadPoolExecutor, TimeoutError
from contextlib import contextmanager, closing
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from storage import Store


class AuditTests(unittest.TestCase):
    """验证公开业务结果，不依赖随机抽取到某个指定学生。"""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = Store(Path(self.temp.name) / "data")
        self.store.mutate("settings", {"values": {"auto_backup": False}})
        self.cid = self.store.mutate("save_class", {"name": "审查班"})["classes"][0][
            "id"
        ]

    def test_xlsx_without_dimensions(self):
        """流式生成的合法 XLSX 没有 dimension，仍需完整读取真实行列。"""
        from openpyxl import Workbook

        book = Workbook(write_only=True)
        sheet = book.create_sheet("名单")
        sheet.append(["姓名", "学号"])
        sheet.append(["测试学生", "0001"])
        path = Path(self.temp.name) / "stream.xlsx"
        book.save(path)
        self.assertEqual(
            Store.read_file(path)["名单"], [["姓名", "学号"], ["测试学生", "0001"]]
        )

    def test_xlsx_without_dimensions_rejects_wide_rows(self):
        """缺少尺寸元数据也不能绕过 50 列限制。"""
        from openpyxl import Workbook

        book = Workbook(write_only=True)
        book.create_sheet("名单").append(["值"] * 51)
        path = Path(self.temp.name) / "wide.xlsx"
        book.save(path)
        with self.assertRaises(ValueError):
            Store.read_file(path)

    def test_xlsx_underreported_dimensions_do_not_truncate_students(self):
        """第三方导出的尺寸声明过小时，仍按实际单元格读取，不静默漏掉学生。"""
        from openpyxl import Workbook
        import re

        path = Path(self.temp.name) / "small-dimension.xlsx"
        book = Workbook()
        book.active.append(["姓名", "学号"])
        book.active.append(["测试学生", "0001"])
        book.save(path)
        with zipfile.ZipFile(path) as archive:
            contents = {
                entry.filename: archive.read(entry.filename)
                for entry in archive.infolist()
            }
        name = "xl/worksheets/sheet1.xml"
        contents[name] = re.sub(
            rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', contents[name]
        )
        with zipfile.ZipFile(path, "w") as archive:
            for name, content in contents.items():
                archive.writestr(name, content)
        self.assertEqual(Store.read_file(path)["Sheet"][1], ["测试学生", "0001"])

    def test_invalid_mutations_do_not_change_data(self):
        """空姓名、布尔 ID、漏传班级及字符串开关不得变成成功的写入。"""
        before = self.store.snapshot()
        for action, payload in [
            ("save_class", {"name": None}),
            ("save_class", {"name": ["班级"]}),
            ("delete_class", {"class_id": True}),
            ("reset", {}),
            ("rules", {"values": {}}),
            ("reset", {"class_id": self.cid, "restore_excluded": "false"}),
            (
                "import",
                {"class_id": self.cid, "rows": [{"name": "甲"}], "replace": "false"},
            ),
        ]:
            with self.subTest(action=action, payload=payload):
                with self.assertRaises(ValueError):
                    self.store.mutate(action, payload)
                self.assertEqual(self.store.snapshot(), before)

    def test_restore_rejects_invalid_class_name(self):
        """物理完整的备份仍可能包含页面无法安全使用的班级名称。"""
        before = self.store.snapshot()
        path = Path(self.store.backup(Path(self.temp.name) / "invalid-name.db"))
        with closing(sqlite3.connect(path)) as db, db:
            db.execute("UPDATE classroom SET name=x'0102'")
        with self.assertRaises(ValueError):
            self.store.restore(path)
        self.assertEqual(self.store.snapshot(), before)

    def test_restore_rejects_unexpected_database_triggers(self):
        """备份不能注入后续名单写入时执行的额外业务触发器。"""
        before = self.store.snapshot()
        path = Path(self.store.backup(Path(self.temp.name) / "trigger.db"))
        with closing(sqlite3.connect(path)) as db, db:
            db.execute(
                "CREATE TRIGGER unexpected AFTER INSERT ON student BEGIN DELETE FROM classroom; END"
            )
        with self.assertRaises(ValueError):
            self.store.restore(path)
        self.assertEqual(self.store.snapshot(), before)

    @unittest.skipUnless(sys.platform == "win32", "真实互斥量测试仅适用于 Windows")
    def test_process_lock_blocks_write_until_restore_scope_exits(self):
        """独立进程持有恢复范围锁时写入必须等待；释放后继续且没有句柄泄漏。"""
        marker = Path(self.temp.name) / "locked"
        script = """import sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
from storage import Store
store = Store(sys.argv[2])
with store.lock:
    Path(sys.argv[3]).write_text('ready', encoding='utf-8')
    sys.stdin.readline()
"""
        child = subprocess.Popen(
            [
                sys.executable,
                "-c",
                script,
                str(Path(__file__).resolve().parents[1] / "backend"),
                str(self.store.folder),
                str(marker),
            ],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
        try:
            deadline = time.monotonic() + 10
            while (
                not marker.exists()
                and time.monotonic() < deadline
                and child.poll() is None
            ):
                time.sleep(0.02)
            self.assertTrue(marker.exists(), "子进程未取得数据目录锁")
            with ThreadPoolExecutor(max_workers=1) as pool:
                pending = pool.submit(
                    self.store.mutate, "save_class", {"name": "等待写入班"}
                )
                try:
                    with self.assertRaises(TimeoutError):
                        pending.result(timeout=0.15)
                finally:
                    child.stdin.write("\n")
                    child.stdin.flush()
                result = pending.result(timeout=5)
                self.assertTrue(
                    any(c["name"] == "等待写入班" for c in result["classes"])
                )
            child.communicate(timeout=5)
            self.assertEqual(child.returncode, 0)
        finally:
            if child.poll() is None:
                child.kill()
                child.communicate(timeout=5)
            for stream in (child.stdin, child.stdout, child.stderr):
                stream.close()

    def test_snapshot_query_count_does_not_grow_with_classes(self):
        """完整快照的读取次数固定，避免每个班级额外发两次查询。"""
        for index in range(6):
            self.store.mutate("save_class", {"name": f"班级{index}"})
        original = self.store.connect
        queries = []

        @contextmanager
        def traced():
            with original() as db:
                db.set_trace_callback(queries.append)
                yield db

        self.store.connect = traced
        self.assertEqual(len(self.store.snapshot()["classes"]), 7)
        self.assertLessEqual(
            sum(q.lstrip().upper().startswith("SELECT") for q in queries), 5
        )

    def test_history_order_uses_index(self):
        """按班级和 ID 倒序分页不应为历史排序构建临时 B 树。"""
        with self.store.connect() as db:
            plan = db.execute(
                "EXPLAIN QUERY PLAN SELECT time,results FROM draw_record "
                "WHERE class_id=? ORDER BY id DESC LIMIT 50",
                (self.cid,),
            ).fetchall()
        self.assertFalse(any("TEMP B-TREE" in row[3] for row in plan), plan)

    def test_snapshot_history_limit_and_class_isolation(self):
        """批量查询仍按班各保留最近 50 次，不串班、不改变快照协议。"""
        import json

        other = self.store.mutate("save_class", {"name": "其他班"})["classes"][-1]["id"]
        with self.store.connect() as db:
            for cid in (self.cid, other):
                db.executemany(
                    "INSERT INTO draw_record(class_id,request,time,results) VALUES(?,?,?,?)",
                    [
                        (
                            cid,
                            f"{cid}-{i}",
                            "2026-09-20T00:00:00+00:00",
                            json.dumps(
                                [{"id": i + 1, "name": f"{cid}-{i}", "no": None}]
                            ),
                        )
                        for i in range(55)
                    ],
                )
        for classroom in self.store.snapshot()["classes"]:
            self.assertEqual(len(classroom["history"]), 50)
            self.assertEqual(
                [r["results"][0]["name"] for r in classroom["history"]],
                [f"{classroom['id']}-{i}" for i in range(54, 4, -1)],
            )

    def test_independent_stores_do_not_lose_order_updates(self):
        """同一学生连续两次上移应移动两位，即使请求来自两个独立 Store。"""
        students = self.store.mutate(
            "add_students",
            {
                "class_id": self.cid,
                "rows": [{"name": name} for name in ["甲", "乙", "丙"]],
            },
        )["classes"][0]["students"]
        sid = students[-1]["id"]
        stores = [Store(self.store.folder), Store(self.store.folder)]

        class DelayedConnection:
            def __init__(self, connection):
                self.connection = connection

            def __getattr__(self, name):
                return getattr(self.connection, name)

            def execute(self, sql, *args):
                cursor = self.connection.execute(sql, *args)
                if sql.strip().startswith("SELECT id FROM student WHERE class_id="):
                    rows = cursor.fetchall()
                    time.sleep(0.1)
                    return rows
                return cursor

        def delayed_connect(original):
            @contextmanager
            def connect():
                with original() as db:
                    yield DelayedConnection(db)

            return connect

        for store in stores:
            store.connect = delayed_connect(store.connect)
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(
                pool.map(
                    lambda store: store.mutate(
                        "move_student",
                        {"class_id": self.cid, "id": sid, "direction": -1},
                    ),
                    stores,
                )
            )
        self.assertEqual(self.store.snapshot()["classes"][0]["students"][0]["id"], sid)


if __name__ == "__main__":
    unittest.main()
