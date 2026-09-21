"""名单下拉及关于页的真实 WebView2 专项回归。

@author ahui
直接运行本脚本；使用新建的隔离数据库，不接触正式用户名单。
指针测试补齐 pointerdown、焦点转移和 click 顺序，避免仅调用 click 漏掉关闭逻辑。
报告及截图保存在 work 下自动创建的回归目录，失败以非零退出码报告。
"""
import json
import os
from pathlib import Path
import sys
import tempfile
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
ARTIFACTS = Path(tempfile.mkdtemp(prefix="roster-about-", dir=ROOT.parent))
os.environ["CLASS_PICKER_DATA_DIR"] = str(ARTIFACTS / "data")

import main
import verify_window
from storage import Store


def run(api):
    """验证关闭/展开两种初态、筛选保持、弹窗焦点、写入忙碌及响应式布局。"""
    window = api._window
    checks = []
    report = {"passed": False, "checks": checks}

    def wait(code, label):
        """等待可观察的 DOM 条件，超时保留失败步骤。"""
        deadline = time.monotonic() + 12
        while time.monotonic() < deadline:
            if window.evaluate_js(code):
                checks.append(label)
                return
            time.sleep(0.08)
        raise AssertionError(label)

    def press(text=None, selector=None, pointer=True):
        """按文案或选择器点击；pointer=False 模拟键盘激活而不发送指针事件。"""
        target = (f"document.querySelector({json.dumps(selector)})" if selector else
                  f"[...document.querySelectorAll('button')].find(b=>b.textContent.trim()==={json.dumps(text)})")
        result = window.evaluate_js(f"""(() => {{
            const button = {target};
            if (!button || button.disabled) return false;
            if ({str(pointer).lower()}) {{
                button.dispatchEvent(new PointerEvent('pointerdown', {{bubbles:true}}));
                button.dispatchEvent(new MouseEvent('mousedown', {{bubbles:true}}));
            }}
            button.focus();
            button.click();
            return true;
        }})()""")
        assert result, text or selector

    def screenshot(name):
        """通过 WebView2 自身截图，仅捕获当前应用内容。"""
        # 等待界面短过渡完成，避免截到旧、新状态之间的颜色。
        time.sleep(0.2)
        from System import Func, Object
        from System.IO import FileStream, FileMode
        from Microsoft.Web.WebView2.Core import CoreWebView2CapturePreviewImageFormat
        stream = FileStream(str(ARTIFACTS / name), FileMode.Create)
        try:
            task = window.native.Invoke(Func[Object](lambda:
                window.native.browser.webview.CoreWebView2.CapturePreviewAsync(
                    CoreWebView2CapturePreviewImageFormat.Png, stream)))
            task.Wait()
        finally:
            stream.Close()

    def fill(selector, value):
        """通过输入事件填写控件，覆盖 Vue 双向绑定而非直接修改组件状态。"""
        window.evaluate_js(f"""(() => {{
            const input = document.querySelector({json.dumps(selector)});
            input.value = {json.dumps(value)};
            input.dispatchEvent(new Event('input', {{bubbles:true}}));
        }})()""")

    try:
        wait("!!document.querySelector('.gear-button') && !!document.querySelector('.class-trigger')", "应用及隔离名单加载")
        assert window.evaluate_js("document.querySelector('.draw-btn').textContent.trim()==='添加学生' && !document.querySelector('.draw-btn').disabled"), "空班级主按钮可直接添加学生"
        assert window.evaluate_js("!document.body.textContent.includes('前往班级名单添加学生')")
        press(selector='.draw-btn')
        wait("!!document.querySelector('.roster-page')", "空班级主按钮进入名单")
        press('返回课堂')
        wait("!!document.querySelector('.recent')", "返回课堂记录入口")
        press(selector='.recent')
        wait("!!document.querySelector('.history-page') && document.querySelector('.history-empty')?.textContent.includes('暂无抽取记录')", "记录页空状态")
        press('返回课堂')
        wait("!!document.querySelector('.gear-button')", "记录页返回课堂")
        press(selector=".gear-button")
        wait("!!document.querySelector('.settings-popover')", "设置菜单展开")
        press("班级名单")
        wait("!!document.querySelector('.roster-page')", "进入班级名单")
        # 两个页面在同一窗口宽度下比较，不能把窄窗口字号与桌面默认字号比较。
        heading_styles = {}
        for width, height in [(1280, 820), (1100, 760), (1000, 700), (800, 560), (1600, 900)]:
            window.resize(width, height)
            time.sleep(0.2)
            heading_styles[width] = window.evaluate_js("""(() => {
                const style = getComputedStyle(document.querySelector('.page-heading h1'));
                return [style.fontSize, style.fontWeight, style.color];
            })()""")
        window.resize(1280, 820)
        opened = "document.querySelector('.class-trigger').getAttribute('aria-expanded')==='true'"
        for pointer in [True, False]:
            for initially_open in [False, True]:
                for label in ["新建班级", "本班规则", "修改名称", "上移", "下移", "删除班级"]:
                    if window.evaluate_js(opened):
                        press(selector=".class-trigger")
                    if initially_open:
                        press(selector=".class-trigger")
                        wait(opened, "准备已展开状态")
                        window.evaluate_js("""(() => {
                            const input = document.querySelector('.class-search input');
                            input.value = '回归';
                            input.dispatchEvent(new Event('input', {bubbles:true}));
                        })()""")
                    # 记录属性变化，已展开时即便最终重开，中间收起也判为失败。
                    window.evaluate_js("""(() => {
                        window.dropdownChanges = [];
                        window.dropdownObserver = new MutationObserver(records => {
                            records.forEach(r => window.dropdownChanges.push(r.oldValue));
                        });
                        window.dropdownObserver.observe(document.querySelector('.class-trigger'), {
                            attributes:true, attributeFilter:['aria-expanded'], attributeOldValue:true
                        });
                    })()""")
                    press(label, pointer=pointer)
                    wait(opened, f"{label}展开：pointer={pointer}, initially_open={initially_open}")
                    if label in ["上移", "下移"]:
                        wait("!document.querySelector('.push-actions button').disabled", "排序提交完成")
                        # 排序后不只检查菜单是否展开，还要检查候选项、勾选与真实班级一致。
                        assert window.evaluate_js("""(() => {
                            const selected = document.querySelectorAll('.class-option[aria-selected="true"]');
                            const input = document.querySelector('.class-search input');
                            const active = document.getElementById(input.getAttribute('aria-activedescendant'));
                            return selected.length === 1 && selected[0] === active
                                && !document.querySelector('.option-active')
                                && selected[0].textContent.trim() === document.querySelector('.class-trigger span').textContent.trim();
                        })()"""), "排序后高亮未跟随当前班级"
                    else:
                        wait("!!document.querySelector('dialog[open]')", "班级弹窗打开")
                        assert window.evaluate_js("document.querySelector('dialog').contains(document.activeElement)"), "下拉抢走弹窗焦点"
                        press("取消")
                        wait("!document.querySelector('dialog[open]')", "关闭班级弹窗")
                    assert window.evaluate_js(opened), f"{label}完成后收起"
                    if initially_open:
                        assert window.evaluate_js("window.dropdownChanges.length===0"), f"{label}发生先关后开"
                        assert window.evaluate_js("document.querySelector('.class-search input').value==='回归'"), "筛选丢失"
                    window.evaluate_js("window.dropdownObserver.disconnect()")
        screenshot("roster-expanded.png")
        # 鼠标候选不产生第二个选中样式；键盘候选仍可见且 Enter 可提交。
        window.evaluate_js("document.querySelector('.class-option:not([aria-selected=true])').dispatchEvent(new PointerEvent('pointermove',{bubbles:true}))")
        assert window.evaluate_js("!document.querySelector('.option-active') && document.querySelectorAll('.class-option[aria-selected=true]').length===1")
        window.evaluate_js("document.querySelector('.class-search input').dispatchEvent(new KeyboardEvent('keydown',{key:'Home',bubbles:true}))")
        wait("!!document.querySelector('.option-active')", "键盘候选可见")
        assert window.evaluate_js("""(() => {
            const candidate = document.querySelector('.option-active');
            return getComputedStyle(candidate).outlineStyle === 'solid'
                && (candidate.getAttribute('aria-selected') === 'true'
                    || getComputedStyle(candidate).backgroundColor === 'rgba(0, 0, 0, 0)');
        })()"""), "候选焦点不应使用选中底色"
        window.evaluate_js("document.querySelector('.class-search input').dispatchEvent(new KeyboardEvent('keydown',{key:'Enter',bubbles:true}))")
        wait(f"!({opened})", "键盘选中后关闭")
        press("本班规则")
        wait("!!document.querySelector('dialog[open]')", "保存提示测试")
        press(selector='dialog .check-field input')
        fill('dialog [aria-label=每次抽取人数]', '2')
        press(selector='dialog .check-field input')
        assert window.evaluate_js("document.querySelector('dialog [aria-label=每次抽取人数]').value==='1'"), '跟随全局必须显示真实全局值'
        press(selector='dialog .check-field input')
        assert window.evaluate_js("document.querySelector('dialog [aria-label=每次抽取人数]').value==='2'"), '取消跟随后保留本班草稿'
        press(selector='dialog .check-field input')
        for width, height in [(1280, 820), (800, 560)]:
            window.resize(width, height)
            time.sleep(0.2)
            assert window.evaluate_js("""[...document.querySelectorAll('.rules-dialog .settings-grid')].every(grid => {
                const rows = [...grid.children];
                return rows.length < 2 || rows[1].getBoundingClientRect().top >= rows[0].getBoundingClientRect().bottom;
            })"""), "规则弹窗必须按自身宽度使用单列"
            screenshot(f'rules-fixed-{width}.png')
        window.resize(1280, 820)
        press("保存")
        wait("!document.querySelector('dialog[open]') && document.querySelector('.notice')?.textContent.trim()==='保存成功'", "本班规则保存成功提示")
        press(selector=".notice button")
        press(selector=".search-field input", pointer=False)
        wait(f"!({opened})", "键盘焦点离开关联操作区正常收起")
        press("下移", pointer=False)
        wait(opened, "键盘操作重新展开")
        wait("!document.querySelector('.push-actions button').disabled", "键盘排序完成")
        press(selector=".search-field input")
        wait(f"!({opened})", "点击无关区域正常收起")
        press(selector=".class-trigger")
        wait(opened, "重新展开")
        window.evaluate_js("document.querySelector('.class-search input').dispatchEvent(new KeyboardEvent('keydown',{key:'Escape',bubbles:true}))")
        wait(f"!({opened})", "Esc 正常关闭")
        press(selector=".class-trigger")
        wait(opened, "选择班级前展开")
        press(selector=".class-option:last-child")
        wait(f"!({opened})", "选中班级正常收起")
        # 以真实文件读取结果预览，替换系统选文件交互，避免人工操作干扰回归。
        import_file = ARTIFACTS / 'import-start-row.csv'
        import_file.write_text('导入说明\n姓名,学号\n张三,202609201940\n李四,202609201941\n王五,202609201942\n', encoding='utf-8-sig')
        original_dialog = window.create_file_dialog
        try:
            window.create_file_dialog = lambda *args, **kwargs: (str(import_file),)
            press('导入名单')
            wait("!!document.querySelector('[aria-label=数据起始行]')", "导入起始行输入框")
        finally:
            window.create_file_dialog = original_dialog
        fill('[aria-label=数据起始行]', '3')
        wait("document.querySelectorAll('.import-preview tbody tr').length===3", "从第三行开始预览")
        assert window.evaluate_js("document.querySelector('[aria-label=学号列]').selectedIndex===0")
        window.evaluate_js("document.querySelector('[aria-label=学号列]').value='2'; document.querySelector('[aria-label=学号列]').dispatchEvent(new Event('change',{bubbles:true}))")
        assert window.evaluate_js("document.querySelector('.import-preview tbody').textContent.includes('202609201940') && !document.querySelector('.import-preview tbody').textContent.includes('导入说明')")
        fill('[aria-label=数据起始行]', '0')
        wait("document.querySelector('.form-error')?.textContent.includes('数据起始行')", "无效起始行校验")
        press('取消')
        wait("!document.querySelector('dialog[open]')", "取消导入")
        # 多表交互复用真实文件桥接：独立映射、切换保留、全选与跨表重复定位。
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = '第一表'
        workbook.active.append(['姓名', '学号'])
        workbook.active.append(['导入甲', '001'])
        other = workbook.create_sheet('第二表')
        other.append(['001', '导入乙'])
        workbook.create_sheet('空表')
        multi_file = ARTIFACTS / 'multi-sheet.xlsx'
        workbook.save(multi_file)
        try:
            window.create_file_dialog = lambda *args, **kwargs: (str(multi_file),)
            press('导入名单')
            wait("document.querySelectorAll('.import-sheet-item').length===3", "识别全部三张工作表")
        finally:
            window.create_file_dialog = original_dialog
        assert window.evaluate_js("document.querySelectorAll('.import-sheet-item input:checked').length===1")
        window.evaluate_js("document.querySelector('[aria-label=学号列]').value='2'; document.querySelector('[aria-label=学号列]').dispatchEvent(new Event('change',{bubbles:true}))")
        press(selector='.import-sheet-item:nth-child(2) input')
        press(selector='.import-sheet-item:nth-child(2) button')
        fill('[aria-label=数据起始行]', '1')
        window.evaluate_js("""for (const [label,value] of [['姓名列','2'],['学号列','1']]) {
            const select=document.querySelector(`[aria-label=${label}]`);
            select.value=value; select.dispatchEvent(new Event('change',{bubbles:true}));
        }""")
        wait("document.querySelector('.import-errors')?.textContent.includes('第二表 · 第 1 行') && document.querySelector('.import-errors')?.textContent.includes('第一表 · 第 2 行')", "跨表重复学号定位两处原行号")
        assert window.evaluate_js("document.querySelector('dialog button[type=submit]').disabled")
        # 清空学号映射后只导入姓名，切换回第一表确认原配置不被覆盖。
        window.evaluate_js("document.querySelector('[aria-label=学号列]').selectedIndex=0; document.querySelector('[aria-label=学号列]').dispatchEvent(new Event('change',{bubbles:true}))")
        wait("!document.querySelector('.import-errors')", "清空学号列解除冲突")
        press(selector='.import-sheet-item:first-child button')
        assert window.evaluate_js("document.querySelector('[aria-label=学号列]').value==='2' && document.querySelector('[aria-label=数据起始行]').value==='2'")
        press('将此配置应用到其他已选表')
        press(selector='.import-sheet-item:nth-child(2) button')
        assert window.evaluate_js("document.querySelector('[aria-label=姓名列]').value==='1' && document.querySelector('[aria-label=数据起始行]').value==='2'")
        fill('[aria-label=数据起始行]', '1')
        window.evaluate_js("""for (const [label,index] of [['姓名列',2],['学号列',0]]) {
            const select=document.querySelector(`[aria-label=${label}]`);
            select.selectedIndex=index; select.dispatchEvent(new Event('change',{bubbles:true}));
        }""")
        wait("!document.querySelector('.import-errors')", "批量应用配置后仍可单独修正")
        press(selector='.import-sheet-item:first-child button')
        press(selector='.import-select-all input')
        wait("document.querySelector('.import-errors')?.textContent.includes('空表')", "空工作表明确提示")
        press(selector='.import-select-all input')
        wait("document.querySelector('dialog button[type=submit]').disabled", "未选择工作表禁用提交")
        press(selector='.import-sheet-item:first-child input')
        press(selector='.import-sheet-item:nth-child(2) input')
        for width, height in [(1280, 820), (800, 560)]:
            window.resize(width, height)
            time.sleep(0.2)
            assert window.evaluate_js("""(() => {
                const content=document.querySelector('dialog .modal-content');
                return content.scrollWidth <= content.clientWidth + 1 &&
                    [...document.querySelectorAll('.import-mapping select')].every(el=>el.getBoundingClientRect().width>80);
            })()"""), '导入弹窗窄屏不得横向溢出'
            screenshot(f'import-multi-{width}.png')
        window.resize(1280, 820)
        # 查看未勾选空表：起始行无效不能通过浏览器原生校验误拦截已选表提交。
        press(selector='.import-sheet-item:nth-child(3) button')
        press('确认导入')
        wait("!document.querySelector('dialog[open]') && document.querySelectorAll('.roster-panel tbody tr').length===2", "多张工作表统一提交")
        imported = [s for c in api._store.snapshot()['classes'] for s in c['students']]
        assert [(s['name'], s['no']) for s in imported] == [('导入甲', '001'), ('导入乙', None)]
        # 通过现有 UI 清理本用例，保持后续草稿回归的空班级前提。
        press(selector='.roster-panel thead input[type=checkbox]')
        press('删除选中（2）')
        wait("document.querySelectorAll('.roster-panel tbody tr').length===0", "清理多表导入测试名单")
        # 完整预览必须覆盖最后一页；错误位于第三页，也必须在首页阻止提交。
        paged_file = ARTIFACTS / 'paged-import.csv'
        paged_file.write_text('姓名,学号\n' + '\n'.join(
            f'{"" if index == 111 else f"分页测试{index}"},{index:04d}'
            for index in range(1, 122)), encoding='utf-8-sig')
        try:
            window.create_file_dialog = lambda *args, **kwargs: (str(paged_file),)
            press('导入名单')
            wait("document.querySelectorAll('.import-preview tbody tr').length===50", "首页只渲染50条，完整数据可分页")
        finally:
            window.create_file_dialog = original_dialog
        assert window.evaluate_js("document.querySelector('dialog button[type=submit]').disabled && document.querySelector('.import-pagination').textContent.includes('121')"), '未显示页的错误仍阻止整批导入'
        press(selector='button[aria-label="下一页"]')
        wait("document.querySelector('.import-preview tbody tr').dataset.row==='52'", "第二页保留原始行号")
        press(selector='button[aria-label="下一页"]')
        wait("document.querySelectorAll('.import-preview tbody tr').length===21 && document.querySelector('button[aria-label=下一页]').disabled", "最后一页完整显示剩余21条")
        press(selector='button[aria-label="上一页"]')
        press(selector='button[aria-label="上一页"]')
        press('查看问题')
        wait("document.querySelector('[aria-label=预览页码]').value==='3' && document.activeElement.dataset.row==='112'", "查看问题跳转第三页并聚焦错误原始行")
        assert window.evaluate_js("document.activeElement.classList.contains('import-row-error')"), '错误行高亮'
        # 行内修正不写源文件；取消保留原值，保存后完整校验更新且留在当前页。
        press(selector='tr[data-row="112"] button[aria-label="编辑预览行"]')
        fill('[aria-label=预览姓名]', '修正测试')
        assert window.evaluate_js("document.querySelector('dialog button[type=submit]').disabled && document.querySelector('[aria-label=姓名列]').matches(':disabled')"), '草稿期间冻结导入与映射'
        press(selector='button[aria-label="取消预览编辑"]')
        assert window.evaluate_js("document.querySelector('tr[data-row=\"112\"]').textContent.includes('缺少姓名')")
        press(selector='tr[data-row="112"] button[aria-label="编辑预览行"]')
        press(selector='button[aria-label="保存预览行"]')
        wait("document.querySelector('[aria-label=预览姓名]') && document.querySelector('.import-row-message[role=alert]')", "空姓名拒绝保存，保留草稿")
        fill('[aria-label=预览姓名]', '修正测试')
        fill('[aria-label=预览学号]', '000111')
        press(selector='button[aria-label="保存预览行"]')
        wait("!document.querySelector('.import-errors') && document.querySelector('[aria-label=预览页码]').value==='3'", "修正后错误解除且保持当前页")
        assert window.evaluate_js("document.querySelector('tr[data-row=\"112\"]').textContent.includes('000111')"), '手填学号保留前导零'
        window.evaluate_js("document.querySelector('[aria-label=学号列]').value='2'; document.querySelector('[aria-label=学号列]').dispatchEvent(new Event('change',{bubbles:true}))")
        wait("document.querySelector('[aria-label=预览页码]').value==='1'", "切换学号列重置页码而非修正数据")
        press(selector='button[aria-label="下一页"]')
        press(selector='button[aria-label="下一页"]')
        assert window.evaluate_js("document.querySelector('tr[data-row=\"112\"]').textContent.includes('修正测试') && document.querySelector('tr[data-row=\"112\"]').textContent.includes('000111')"), '切换学号列保留已编辑姓名及学号'
        press(selector='tr[data-row="112"] button[aria-label="移除预览行"]')
        wait("!document.querySelector('tr[data-row=\"112\"]') && document.querySelector('.import-pagination').textContent.includes('120')", "移除更新完整预览人数")
        press('撤销移除（1）')
        wait("document.querySelector('tr[data-row=\"112\"]')?.textContent.includes('修正测试')", "撤销恢复行和修正")
        screenshot('import-row-edited.png')
        # 用新读入文件验证修正不污染源文件，后续继续覆盖原错误分页布局。
        press('取消')
        try:
            window.create_file_dialog = lambda *args, **kwargs: (str(paged_file),)
            press('导入名单')
            wait("!!document.querySelector('.import-errors')", "重新打开源文件仍保留原始问题")
        finally:
            window.create_file_dialog = original_dialog
        press('查看问题')
        wait("document.querySelector('[aria-label=预览页码]').value==='3'", "重新定位原错误")
        for width, height in [(1280, 820), (800, 560)]:
            window.resize(width, height)
            time.sleep(0.2)
            assert window.evaluate_js("""(() => {
                const content = document.querySelector('dialog .modal-content');
                const preview = document.querySelector('.import-preview').getBoundingClientRect();
                const pager = document.querySelector('.import-pagination').getBoundingClientRect();
                const footer = document.querySelector('dialog footer').getBoundingClientRect();
                return content.scrollHeight <= content.clientHeight + 1
                    && content.scrollWidth <= content.clientWidth + 1
                    && preview.height >= 60 && preview.bottom <= pager.top + 1
                    && pager.bottom <= footer.top + 1;
            })()"""), '分页、表格与操作栏不遮挡，外层无滚动'
            screenshot(f'import-error-paged-{width}.png')
        window.resize(1280, 820)
        fill('[aria-label=数据起始行]', '113')
        wait("document.querySelector('[aria-label=预览页码]').value==='1' && document.querySelectorAll('.import-preview tbody tr').length===10", "调整范围重置分页，末尾10条可完整预览")
        screenshot('import-paged-valid.png')
        press('取消')
        wait("!document.querySelector('dialog[open]')", "分页测试取消，不写入名单")
        # 多行录入：重复追加保留输入；缺姓名、重复学号均不产生部分保存。
        draft_class = window.evaluate_js("document.querySelector('.class-trigger span').textContent.trim()")
        assert window.evaluate_js("!document.querySelector('input[aria-label=新增行数]')"), "不应展示行数输入框"
        for _ in range(3):
            press("添加学生")
        wait("document.querySelectorAll('.new-student-row').length===3", "连续添加三行")
        assert window.evaluate_js("document.activeElement.id.startsWith('new-name-') && !document.querySelector('dialog[open]')")
        fill('[aria-label="新增第1行姓名"]', '新增甲')
        fill('[aria-label="新增第2行姓名"]', '新增乙')
        fill('[aria-label="新增第2行学号"]', '1')
        press("添加学生")
        press("添加学生")
        wait("document.querySelectorAll('.new-student-row').length===5", "继续追加两行")
        assert window.evaluate_js("document.querySelector('[aria-label=新增第1行姓名]').value==='新增甲'")
        fill('[aria-label="新增第4行姓名"]', '新增丙')
        fill('[aria-label="新增第4行学号"]', '0008')
        press(selector='.new-student-row:nth-child(4) input[type="checkbox"]')
        fill('[aria-label="新增第5行学号"]', '1')
        press("保存新增")
        wait("document.querySelector('.draft-row-error')?.textContent.includes('请填写姓名')", "定位缺少姓名行")
        assert sum(len(c['students']) for c in api._store.snapshot()['classes']) == 0
        fill('[aria-label="新增第5行姓名"]', '新增丁')
        press("保存新增")
        wait("document.querySelector('.draft-row-error')?.textContent.includes('学号重复')", "批内重复学号阻止保存")
        fill('[aria-label="新增第5行学号"]', '3')
        press(selector='.notice button')
        for width, height in [(1280, 820), (800, 560)]:
            window.resize(width, height)
            time.sleep(0.2)
            # 聚焦末行再返回首行，验证固定表头不会遮住键盘/校验定位的输入框。
            for row_index in [5, 1]:
                window.evaluate_js(f"document.querySelector('[aria-label=新增第{row_index}行姓名]').focus()")
                assert window.evaluate_js("""(() => {
                    const input = document.activeElement.getBoundingClientRect();
                    const header = document.querySelector('.table-scroll th').getBoundingClientRect();
                    const table = document.querySelector('.table-scroll').getBoundingClientRect();
                    return input.top >= header.bottom && input.bottom <= table.bottom;
                })()"""), "固定表头不得遮挡焦点输入框"
            window.evaluate_js("document.querySelector('.table-scroll').scrollTop=0")
            screenshot(f"roster-drafts-{width}.png")
            draft_metrics = window.evaluate_js("""({
                width:innerWidth, scrollWidth:document.documentElement.scrollWidth,
                tableHeight:document.querySelector('.table-scroll').clientHeight,
                toolbarHeight:document.querySelector('.roster-toolbar').clientHeight,
                draftToolbarHeight:document.querySelector('.draft-toolbar').clientHeight
            })""")
            assert draft_metrics['scrollWidth'] <= draft_metrics['width'] and draft_metrics['tableHeight'] > 80, draft_metrics
        window.resize(1280, 820)
        press(selector='.class-trigger')
        wait(opened, "切换班级前展开")
        other_class = window.evaluate_js("document.querySelector('.class-option:not([aria-selected=true])').textContent.trim()")
        press(selector='.class-option:not([aria-selected=true])')
        wait("document.querySelectorAll('.new-student-row').length===0", "另一班级草稿独立")
        press("添加学生")
        press("添加学生")
        wait("document.querySelectorAll('.new-student-row').length===2", "另一班追加草稿")
        fill('[aria-label="新增第1行姓名"]', '其他班草稿')
        press(selector='.class-trigger')
        wait(opened, "恢复原班级")
        window.evaluate_js(f"[...document.querySelectorAll('.class-option')].find(e=>e.textContent.trim()==={json.dumps(draft_class)}).click()")
        wait("document.querySelectorAll('.new-student-row').length===5", "切班后恢复原有输入")
        assert window.evaluate_js("document.querySelector('[aria-label=新增第1行姓名]').value==='新增甲'")
        press("返回课堂")
        wait("!!document.querySelector('dialog[open]')", "返回前保护未保存新增")
        press("取消")
        wait("!document.querySelector('dialog[open]')", "继续填写")
        press(selector='button[aria-label="关闭"]')
        wait("!!document.querySelector('dialog[open]')", "关闭按钮保护未保存新增")
        press("取消")
        wait("!document.querySelector('dialog[open]')", "取消关闭")
        window.evaluate_js("""(() => {
            const save = [...document.querySelectorAll('.draft-toolbar button')].find(b=>b.textContent==='保存新增');
            save.click(); save.click();
        })()""")
        wait("!document.querySelector('.new-student-row') && document.querySelector('.notice')?.textContent.trim()==='保存成功'", "批量提交成功")
        saved = next(c['students'] for c in api._store.snapshot()['classes'] if c['name'] == draft_class)
        assert len(saved) == 4, "空行未忽略或连续点击重复新增"
        by_name = {student['name']: student for student in saved}
        assert by_name['新增甲']['no'] is None and by_name['新增丙']['no'] == '0008'
        assert by_name['新增丙']['enabled'] == 0
        press(selector='[aria-label=编辑学生]')
        wait("!!document.querySelector('[aria-label=编辑姓名]')", '已存学生行编辑')
        fill('[aria-label=编辑姓名]', '未保存修改')
        press('返回课堂')
        wait("!!document.querySelector('dialog[open]')", '返回保护行编辑')
        press('取消')
        wait("!document.querySelector('dialog[open]')", '取消后保留行编辑')
        assert window.evaluate_js("document.querySelector('[aria-label=编辑姓名]').value==='未保存修改'")
        press(selector='tbody tr:nth-child(2) [aria-label=编辑学生]')
        wait("document.querySelector('dialog h2')?.textContent.includes('放弃本行')", '切行保护修改')
        press('放弃修改')
        wait("!document.querySelector('dialog[open]') && document.querySelector('[aria-label=编辑姓名]').value==='新增乙'", '确认后切换编辑行')
        press('保存')
        wait("!document.querySelector('[aria-label=编辑姓名]')", '结束行编辑')
        # 系统文件框替身计数，连续 click 必须只调用一次下载桥接。
        template_calls = []
        original_dialog = window.create_file_dialog
        def template_dialog(*args, **kwargs):
            template_calls.append(1)
            time.sleep(0.2)
            return str(ARTIFACTS / 'download-template.xlsx')
        try:
            window.create_file_dialog = template_dialog
            window.evaluate_js("""(() => {
                const button=[...document.querySelectorAll('button')].find(b=>b.textContent.trim()==='下载导入模板');
                button.click();button.click();
            })()""")
            wait("document.querySelector('.notice')?.textContent.trim()==='导入模板已下载'", '模板下载完成')
            assert len(template_calls) == 1, '连续点击重复打开保存框'
        finally:
            window.create_file_dialog = original_dialog
        # 已保存行可上下移动；排序反馈直接体现在表格，首尾按钮禁用。
        assert window.evaluate_js("document.querySelector('[aria-label=上移学生]').disabled")
        press(selector='tbody tr:nth-child(2) [aria-label=上移学生]')
        wait("document.querySelector('tbody tr').textContent.includes('新增乙') && !document.querySelector('tbody tr [aria-label=下移学生]').disabled", "学生上移保存")
        press(selector='tbody tr:first-child [aria-label=下移学生]')
        wait("document.querySelector('tbody tr').textContent.includes('新增甲')", "学生下移保存")
        for width, height in [(1280, 820), (800, 560)]:
            window.resize(width, height)
            time.sleep(0.2)
            assert window.evaluate_js("document.querySelector('.table-scroll').scrollWidth <= document.querySelector('.table-scroll').clientWidth"), "四个操作按钮不得造成横向溢出"
            screenshot(f'roster-actions-{width}.png')
        window.resize(1280, 820)
        press(selector='.class-trigger')
        wait(opened, "检查另一班草稿")
        window.evaluate_js(f"[...document.querySelectorAll('.class-option')].find(e=>e.textContent.trim()==={json.dumps(other_class)}).click()")
        wait("document.querySelectorAll('.new-student-row').length===2", "保存本班不清除其他班草稿")
        press("取消新增")
        wait("!!document.querySelector('dialog[open]')", "取消填写内容需要确认")
        press("放弃新增")
        wait("!document.querySelector('.new-student-row') && !document.querySelector('dialog[open]')", "仅清除本班新增草稿")
        press("添加学生")
        press("添加学生")
        wait("document.querySelectorAll('.new-student-row').length===2", "全空行可直接清除")
        assert window.evaluate_js("[...document.querySelectorAll('.draft-toolbar button')].find(b=>b.textContent==='保存新增').disabled")
        press("取消新增")
        wait("!document.querySelector('.new-student-row') && !document.querySelector('dialog[open]')", "全空行取消不弹窗")
        assert window.evaluate_js("!document.querySelector('.roster-page').textContent.includes('编号')"), "名单页应使用学号文案"
        press("返回课堂")
        wait("!!document.querySelector('.draw-btn')", "返回后检查空学号抽取结果")
        press(selector='.class-trigger')
        wait(opened, "选择多行录入的班级")
        window.evaluate_js(f"[...document.querySelectorAll('.class-option')].find(e=>e.textContent.trim()==={json.dumps(draft_class)}).click()")
        wait(f"document.querySelector('.class-trigger span').textContent.trim()==={json.dumps(draft_class)}", "课堂切换完成")
        blank_number_checked = False
        for _ in range(3):
            # 后端提交不等于前端揭晓；动画中记录与统计保持上一轮，学号隐藏。
            previous_recent = window.evaluate_js("document.querySelector('.recent').textContent")
            previous_stats = window.evaluate_js("document.querySelector('.stats').textContent")
            press(selector='.draw-btn')
            wait("document.querySelector('.draw-btn').disabled", "开始抽取动画")
            assert window.evaluate_js("document.querySelector('.recent').textContent") == previous_recent
            assert window.evaluate_js("document.querySelector('.stats').textContent") == previous_stats
            assert window.evaluate_js("document.querySelector('.number').textContent.trim()===''"), "动画期间不得显示旧学号"
            time.sleep(0.3)
            assert window.evaluate_js("document.querySelector('.recent').textContent") == previous_recent
            wait("!document.querySelector('.draw-btn').disabled && document.querySelector('.name').textContent.includes('新增')", "录入名单可正常抽取")
            latest = next(c['latest'] for c in api._store.snapshot()['classes'] if c['name'] == draft_class)[0]
            if latest['no'] is None:
                assert window.evaluate_js("document.querySelector('.number').textContent.trim()===''"), "空学号不能展示 null 或自动号码"
                blank_number_checked = True
        assert blank_number_checked, "空学号学生应正常参与抽取"
        for width, height in [(1280, 820), (800, 560)]:
            window.resize(width, height)
            time.sleep(0.2)
            assert window.evaluate_js("""(() => {
                const head=document.querySelector('.recent-head'), item=document.querySelector('.recent-item');
                return getComputedStyle(document.querySelector('.recent')).display==='block'
                    && item.getBoundingClientRect().top >= head.getBoundingClientRect().bottom
                    && head.scrollWidth<=head.clientWidth;
            })()"""), "最近记录保持上下两行且标题不挤压"
            screenshot(f'classroom-fixed-{width}.png')
        window.resize(1280, 820)
        press(selector='.recent')
        wait("document.querySelectorAll('.history-record').length===3", "记录页显示抽取明细")
        assert window.evaluate_js("document.querySelector('.history-list').textContent.includes('新增甲') && !document.querySelector('.history-list').textContent.includes('null')")
        for width, height in [(1280, 820), (800, 560)]:
            window.resize(width, height)
            time.sleep(0.2)
            assert window.evaluate_js("document.querySelector('.history-list').clientHeight > 80 && document.documentElement.scrollWidth <= innerWidth"), "记录页可滚动且无横向溢出"
            screenshot(f'history-{width}.png')
        window.resize(1280, 820)
        press(selector='.class-trigger')
        wait(opened, '记录页切换班级')
        press(selector='.class-option:not([aria-selected=true])')
        wait("document.querySelector('.history-empty')?.textContent.includes('暂无抽取记录')", '记录按班级隔离')
        press('返回课堂')
        wait("!!document.querySelector('.draw-btn')", '查看记录后返回课堂')
        press(selector=".gear-button")
        wait("!!document.querySelector('.settings-popover')", "关于入口")
        press("关于")
        wait("!!document.querySelector('.about-page')", "关于页面")
        assert window.evaluate_js("document.querySelectorAll('.about-tabs [role=tab]').length===3 && !document.querySelector('#about-tab-version')"), "仅保留三个内容主题"
        assert window.evaluate_js("document.querySelector('.about-heading button').textContent.trim()==='返回课堂'")
        for width, height in [(1280, 820), (1100, 760), (1000, 700), (800, 560), (1600, 900)]:
            window.resize(width, height)
            time.sleep(0.3)
            metrics = window.evaluate_js("""(() => {
                const panel=document.querySelector('.about-scroll');
                const column=document.querySelector('.about-column');
                return {width:innerWidth, overflow:panel.scrollWidth>panel.clientWidth,
                    rows:[...column.querySelectorAll('li')].map(li=>{
                        const range=document.createRange(); range.selectNodeContents(li);
                        const rect=range.getBoundingClientRect();
                        return {lines:range.getClientRects().length, right:rect.right,
                            boundary:column.getBoundingClientRect().right};
                    })};
            })()""")
            assert not metrics["overflow"], metrics
            assert all(r["lines"] == 1 and r["right"] <= r["boundary"] - 10 for r in metrics["rows"]), metrics
            checks.append({"about_layout": metrics})
            if width in [1280, 800]:
                screenshot(f"about-{width}.png")
            assert window.evaluate_js("""(() => {
                const style = getComputedStyle(document.querySelector('.about-heading h1'));
                return [style.fontSize, style.fontWeight, style.color];
            })()""") == heading_styles[width], "关于页标题应与名单页一致"
            for label, name in [("功能介绍", "features"), ("使用方式", "usage"), ("数据说明", "data")]:
                press(label)
                wait(f"document.querySelector('.about-tabs [aria-selected=true]').textContent.trim()==={json.dumps(label)}", label)
                assert window.evaluate_js("document.querySelector('.about-scroll').scrollWidth<=document.querySelector('.about-scroll').clientWidth")
                if label == "数据说明":
                    assert window.evaluate_js("parseFloat(getComputedStyle(document.querySelector('.about-info-grid')).gap)>=14")
                # 所有主题都只有一个独立版本模块；正文与模块不重叠，窄窗口可滚动查看。
                assert window.evaluate_js("""(() => {
                    const footers = document.querySelectorAll('.about-version-footer');
                    if (footers.length !== 1) return false;
                    const footer = footers[0], scroll = document.querySelector('.about-scroll');
                    const separated = footer.getBoundingClientRect().top
                        >= document.querySelector('.about-main').getBoundingClientRect().bottom + 15;
                    const small = parseFloat(getComputedStyle(footer).fontSize) === 13;
                    scroll.scrollTop = scroll.scrollHeight;
                    const reachable = footer.getBoundingClientRect().bottom <= scroll.getBoundingClientRect().bottom + 1;
                    scroll.scrollTop = 0;
                    return separated && small && reachable && footer.textContent.includes('v1.1.0')
                        && footer.textContent.includes('ahui') && footer.scrollWidth <= footer.clientWidth;
                })()"""), f"{label}底部版本模块布局"
                if width in [1280, 800]:
                    screenshot(f"about-{name}-{width}.png")
            press("功能介绍")
        # 新增 Tab 键盘导航与选择同步验证。
        window.evaluate_js("document.querySelector('[role=tab][aria-selected=true]').dispatchEvent(new KeyboardEvent('keydown',{key:'End',bubbles:true}))")
        wait("document.activeElement.textContent.trim()==='数据说明' && !!document.querySelector('.about-info-grid')", "Tab 键盘导航")
        press("返回课堂")
        wait("!!document.querySelector('.draw-btn')", "关于返回课堂有效")
        press(selector=".gear-button")
        wait("!!document.querySelector('.settings-popover')", "设置入口")
        press("应用设置")
        wait("!!document.querySelector('.settings-page')", "设置保存测试")
        window.evaluate_js("document.querySelector('.stepper input').focus()")
        assert window.evaluate_js("getComputedStyle(document.querySelector('.stepper input')).outlineStyle==='none' && getComputedStyle(document.querySelector('.stepper input')).boxShadow==='none'"), "步进输入框不能出现内部焦点边框"
        window.evaluate_js("""(() => {
            const input = document.querySelector('input[aria-label="每次抽取人数"]');
            input.value = '2';
            input.dispatchEvent(new Event('input',{bubbles:true}));
        })()""")
        press("保存设置")
        wait("document.querySelector('.notice')?.textContent.trim()==='保存成功' && document.querySelector('.save-status').textContent.includes('所有设置已保存')", "应用设置保存成功提示")
        assert api._store.snapshot()["settings"]["count"] == 2
        report["passed"] = True
    except Exception as exc:
        report["error"] = str(exc)
    finally:
        (ARTIFACTS / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        window.destroy()


if __name__ == "__main__":
    store = Store(ARTIFACTS / "data")
    for name in ["回归一班", "回归二班", "回归三班"]:
        store.mutate("save_class", {"name": name})
    verify_window.run = run
    sys.argv.append("--verify-window")
    main.main()
    result = json.loads((ARTIFACTS / "report.json").read_text(encoding="utf-8"))
    print(json.dumps({"passed": result["passed"], "error": result.get("error"), "report": str(ARTIFACTS)}, ensure_ascii=False))
    sys.exit(0 if result["passed"] else 1)
